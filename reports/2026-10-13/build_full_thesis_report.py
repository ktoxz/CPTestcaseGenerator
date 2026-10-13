from __future__ import annotations

import importlib.util
import json
import math
import statistics
from collections import Counter, defaultdict
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
OUT_DIR = ROOT / "reports" / "2026-10-13"
ASSET_DIR = OUT_DIR / "assets"
WORK_DIR = OUT_DIR / ".thesis-work"
OUTPUT = OUT_DIR / "Bao-cao-KLTN-Final-TestcaseGenerator-day-du-2026-10-13-v2.docx"
OLD_BUILDER = ROOT / "reports" / "2026-10-12" / "build_final_progress_report.py"

NAVY = "17365D"
BLUE = "2F5597"
PALE_BLUE = "EAF2F8"
PALE_GRAY = "F5F7F9"
LIGHT_BORDER = "B7C9DB"
WHITE = "FFFFFF"
GREEN = "1B5E20"
AMBER = "8A4B08"
RED = "8B1E1E"


def load_old_module():
    spec = importlib.util.spec_from_file_location("progress_report", OLD_BUILDER)
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    spec.loader.exec_module(module)
    return module


old = load_old_module()


def pct(value, digits=1):
    if value is None:
        return "—"
    return f"{value * 100:.{digits}f}%".replace(".", ",")


def num(value, digits=1):
    if value is None:
        return "—"
    if isinstance(value, int):
        return f"{value:,}".replace(",", ".")
    return f"{value:.{digits}f}".replace(".", ",")


def dataset_of(problem):
    if problem.startswith("p") and problem[1:].isdigit():
        return "CodeNet"
    if problem.lower().startswith("cf") or problem == "G1":
        return "Codeforces"
    return "Khác"


def classify(tpr, tnr):
    if tpr is None or tnr is None:
        return "Chưa đo đủ"
    if tpr >= 0.9 and tnr >= 0.9:
        return "Mạnh"
    if tpr >= 0.8 and tnr >= 0.8:
        return "Đạt 0,8"
    return "Cần cải thiện"


def set_repeat_header(section, text):
    p = section.header.paragraphs[0]
    p.clear()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string("666666")
    p_pr = p._p.get_or_add_pPr()
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "3")
    bottom.set(qn("w:color"), "B7C9DB")
    border.append(bottom)
    p_pr.append(border)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, end])


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=75, start=95, bottom=75, end=95):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=LIGHT_BORDER, size="5"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def configure_styles(doc):
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.0)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing = 1.28
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.first_line_indent = Cm(0.75)
    for name, size, color, before, after in (
        ("Title", 18, NAVY, 0, 12),
        ("Subtitle", 13, "444444", 0, 8),
        ("Heading 1", 16, NAVY, 12, 8),
        ("Heading 2", 14, BLUE, 10, 6),
        ("Heading 3", 12.5, "1F4E79", 8, 4),
    ):
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.first_line_indent = Cm(0)
    doc.styles["Title"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.styles["Subtitle"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for style_name in ("List Bullet", "List Number"):
        st = doc.styles[style_name]
        st.font.name = "Times New Roman"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        st.font.size = Pt(12.5)
        st.paragraph_format.space_after = Pt(3)


def add_table(doc, headers, rows, widths=None, font_size=9, aligns=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    hdr = table.rows[0]
    repeat_table_header(hdr)
    prevent_row_split(hdr)
    for i, text in enumerate(headers):
        cell = hdr.cells[i]
        cell.text = str(text)
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        if widths:
            cell.width = Cm(widths[i])
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.space_after = Pt(0)
            for r in p.runs:
                r.bold = True
                r.font.name = "Times New Roman"
                r.font.size = Pt(font_size)
                r.font.color.rgb = RGBColor.from_string(WHITE)
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        prevent_row_split(table.rows[-1])
        for i, value in enumerate(row):
            cell = cells[i]
            cell.text = str(value)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if row_index % 2 == 1:
                set_cell_shading(cell, PALE_GRAY)
            if widths:
                cell.width = Cm(widths[i])
            for p in cell.paragraphs:
                p.alignment = aligns[i] if aligns else WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.first_line_indent = Cm(0)
                p.paragraph_format.line_spacing = 1.05
                p.paragraph_format.space_after = Pt(0)
                for r in p.runs:
                    r.font.name = "Times New Roman"
                    r.font.size = Pt(font_size)
    set_table_borders(table)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Cm(0.8)
        p.paragraph_format.first_line_indent = Cm(-0.35)
        p.add_run(item)


def add_numbered(doc, items):
    # Write the ordinal explicitly so every logical list restarts at 1. Word's
    # built-in List Number style otherwise continues numbering across chapters.
    for index, item in enumerate(items, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.8)
        p.paragraph_format.first_line_indent = Cm(-0.55)
        p.paragraph_format.space_after = Pt(3)
        p.add_run(f"{index}. ").bold = True
        p.add_run(item)


def add_note(doc, text, title="Ghi chú"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, PALE_BLUE)
    set_cell_margins(cell, 140, 160, 140, 160)
    set_table_borders(table, color="8EA9C1", size="7")
    p = cell.paragraphs[0]
    p.paragraph_format.first_line_indent = Cm(0)
    r = p.add_run(title + ". ")
    r.bold = True
    p.add_run(text)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_caption(doc, number, title, source):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(f"Hình {number}. {title}. ")
    r.bold = True
    r.font.size = Pt(10)
    r2 = p.add_run(f"Nguồn: {source}")
    r2.italic = True
    r2.font.size = Pt(10)


def add_picture(doc, path, width, number, title, source):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Inches(width))
    add_caption(doc, number, title, source)


def add_toc(doc):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    run = p.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\o "1-3" \\h \\z \\u '
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "Mục lục sẽ được cập nhật khi mở trong Microsoft Word."
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, placeholder, end])


def chapter(doc, title):
    doc.add_page_break()
    doc.add_heading(title, level=1)


def create_charts(e, s):
    WORK_DIR.mkdir(parents=True, exist_ok=True)
    charts = old.make_charts(e, s)
    font = ImageFont.truetype("C:/Windows/Fonts/times.ttf", 25)
    small = ImageFont.truetype("C:/Windows/Fonts/times.ttf", 20)
    bold = ImageFont.truetype("C:/Windows/Fonts/timesbd.ttf", 29)

    # Rebuild the source-group chart because G1 is a manually crawled
    # Codeforces problem, even though its identifier does not start with "cf".
    groups = defaultdict(list)
    for run in s["measured"]:
        groups[dataset_of(run["problem"])].append(run)
    labels = [name for name in ("CodeNet", "Codeforces") if groups.get(name)]
    img = Image.new("RGB", (1180, 650), "white")
    d = ImageDraw.Draw(img)
    left, top, right, bottom = 115, 90, 1090, 545
    d.text(((left + right) / 2, 25), "Macro TPR và TNR theo nguồn dữ liệu", font=bold, fill="#111111", anchor="ma")
    for k in range(6):
        value = k / 5
        y = bottom - value * (bottom - top)
        d.line((left, y, right, y), fill="#E2E7EC", width=1)
        d.text((left - 16, y), f"{value:.1f}", font=small, fill="#555555", anchor="rm")
    slot = (right - left) / len(labels)
    width = 125
    for i, label in enumerate(labels):
        center = left + slot * (i + 0.5)
        tpr = statistics.mean(run["oracle_tpr"] for run in groups[label])
        tnr = statistics.mean(run["oracle_tnr"] for run in groups[label])
        for value, offset, color in ((tpr, -width / 2, "#2F5597"), (tnr, width / 2, "#70AD47")):
            x0 = center + offset - width / 2
            x1 = center + offset + width / 2
            y0 = bottom - value * (bottom - top)
            d.rectangle((x0, y0, x1, bottom), fill=color)
            d.text(((x0 + x1) / 2, y0 - 8), f"{value:.3f}", font=small, fill="#222222", anchor="ms")
        d.text((center, bottom + 22), label, font=font, fill="#222222", anchor="ma")
    d.line((left, top, left, bottom), fill="#555555", width=2)
    d.line((left, bottom, right, bottom), fill="#555555", width=2)
    d.rectangle((435, 600, 460, 622), fill="#2F5597")
    d.text((470, 611), "TPR", font=small, fill="#222222", anchor="lm")
    d.rectangle((635, 600, 660, 622), fill="#70AD47")
    d.text((670, 611), "TNR", font=small, fill="#222222", anchor="lm")
    performance = WORK_DIR / "performance-by-dataset-corrected.png"
    img.save(performance)
    charts["performance"] = performance

    # Quality scatter: every point is a completed, locked suite with oracle metrics.
    img = Image.new("RGB", (1180, 760), "white")
    d = ImageDraw.Draw(img)
    left, top, right, bottom = 115, 90, 1090, 650
    d.text(((left + right) / 2, 25), "Phân bố chất lượng theo từng bài", font=bold, fill="#111111", anchor="ma")
    for k in range(6):
        v = 0.5 + k * 0.1
        x = left + (v - 0.3) / 0.7 * (right - left)
        y = bottom - (v - 0.3) / 0.7 * (bottom - top)
        d.line((x, top, x, bottom), fill="#E2E7EC", width=1)
        d.line((left, y, right, y), fill="#E2E7EC", width=1)
        d.text((x, bottom + 16), f"{v:.1f}", font=small, fill="#555555", anchor="ma")
        d.text((left - 15, y), f"{v:.1f}", font=small, fill="#555555", anchor="rm")
    for threshold, color in ((0.8, "#ED7D31"), (0.9, "#70AD47")):
        x = left + (threshold - 0.3) / 0.7 * (right - left)
        y = bottom - (threshold - 0.3) / 0.7 * (bottom - top)
        d.line((x, top, x, bottom), fill=color, width=2)
        d.line((left, y, right, y), fill=color, width=2)
    colors = {"CodeNet": "#2F5597", "Codeforces": "#C00000"}
    for r in s["measured"]:
        x = left + (max(0.3, r["oracle_tnr"]) - 0.3) / 0.7 * (right - left)
        y = bottom - (max(0.3, r["oracle_tpr"]) - 0.3) / 0.7 * (bottom - top)
        d.ellipse((x - 6, y - 6, x + 6, y + 6), fill=colors[dataset_of(r["problem"])], outline="white", width=1)
    d.line((left, top, left, bottom), fill="#555555", width=2)
    d.line((left, bottom, right, bottom), fill="#555555", width=2)
    d.text(((left + right) / 2, 715), "TNR — tỷ lệ chương trình sai bị bắt", font=font, fill="#222222", anchor="ma")
    # Keep the y-axis label compact so it remains inside Word's image crop.
    d.text((68, (top + bottom) / 2), "TPR", font=font, fill="#222222", anchor="mm")
    for i, (label, color) in enumerate(colors.items()):
        x = 390 + i * 210
        d.ellipse((x, 675, x + 15, 690), fill=color)
        d.text((x + 23, 683), label, font=small, fill="#222222", anchor="lm")
    scatter = WORK_DIR / "quality-scatter.png"
    img.save(scatter)
    charts["scatter"] = scatter

    # Failure categories from completed batches only.
    failures = [a for a in e["complete_attempts"] if a["status"] == "failed"]
    counts = Counter(a["failure_reason"] for a in failures)
    labels = [x[0] for x in counts.most_common()]
    values = [counts[x] for x in labels]
    short = {
        "Không sinh đủ AC hoặc validator": "Thiếu AC/validator",
        "Không đủ mutant WA được chứng nhận": "Thiếu mutant WA",
        "Không đủ reference qua biên dịch và sample": "Thiếu reference",
        "Remote judge không sẵn sàng": "Remote judge",
    }
    img = Image.new("RGB", (1180, 620), "white")
    d = ImageDraw.Draw(img)
    d.text((590, 25), "Nguyên nhân thất bại trong các batch hoàn tất", font=bold, fill="#111111", anchor="ma")
    left, top, right, bottom = 390, 100, 1090, 535
    vmax = max(values) if values else 1
    for i, (label, value) in enumerate(zip(labels, values)):
        y0 = top + i * 95
        y1 = y0 + 50
        d.text((left - 20, (y0 + y1) / 2), short.get(label, label), font=small, fill="#222222", anchor="rm")
        x1 = left + value / vmax * (right - left)
        d.rectangle((left, y0, x1, y1), fill="#5B9BD5")
        d.text((x1 + 12, (y0 + y1) / 2), str(value), font=small, fill="#222222", anchor="lm")
    failures_chart = WORK_DIR / "failure-categories.png"
    img.save(failures_chart)
    charts["failures"] = failures_chart
    return charts


def p(doc, text, bold_lead=None):
    para = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        para.add_run(bold_lead).bold = True
        para.add_run(text[len(bold_lead):])
    else:
        para.add_run(text)
    return para


def build_document(e, s, charts):
    doc = Document()
    configure_styles(doc)
    doc.settings.odd_and_even_pages_header_footer = False
    section = doc.sections[0]
    section.different_first_page_header_footer = False
    set_repeat_header(section, "Khóa luận tốt nghiệp — Final Testcase Generator")
    add_page_number(section.footer.paragraphs[0])

    title = doc.add_paragraph(style="Title")
    title.add_run("HỆ THỐNG SINH VÀ CHỨNG NHẬN TESTCASE TỰ ĐỘNG\nCHO BÀI TOÁN LẬP TRÌNH")
    sub = doc.add_paragraph(style="Subtitle")
    sub.add_run("Báo cáo khóa luận tốt nghiệp — bản nội dung không gồm bìa")
    sub = doc.add_paragraph(style="Subtitle")
    sub.add_run("Final Testcase Generator · Chốt dữ liệu ngày 13 tháng 10 năm 2026")

    doc.add_heading("TÓM TẮT", level=1)
    p(doc,
      "Khóa luận nghiên cứu một pipeline dùng mô hình ngôn ngữ lớn để phân tích đề lập trình, sinh chương trình tham chiếu và chương trình sai, "
      "tạo testcase theo các nghĩa vụ phủ, chứng nhận expected output bằng đồng thuận nhiều lời giải, khóa bộ test và đánh giá độc lập trên nguồn chương trình lịch sử. "
      "Mục tiêu không chỉ là tạo input hợp lệ mà còn đồng thời bảo toàn chương trình đúng và phát hiện chương trình sai. Hai đại lượng trung tâm là TPR, tỷ lệ chương trình mang nhãn đúng vẫn được chấp nhận, và TNR, tỷ lệ chương trình mang nhãn sai bị bộ test phát hiện.")
    p(doc,
      f"Tại thời điểm chốt, kho dữ liệu có {s['problem_dirs']} thư mục bài toán, {s['oracle_bundle_count']} oracle bundle và {num(s['manifest_unique'])} mã nguồn phân biệt theo SHA-256. "
      f"Năm batch hoàn tất gồm {s['complete_batch_attempts']} lượt bài; {s['complete_batch_success']} lượt thành công và {s['complete_batch_failed']} lượt thất bại. "
      f"Sau khi lấy report thành công gần nhất theo từng bài, hệ thống có {s['ever_success_count']} bộ test đã khóa, tổng cộng {num(s['cases_total'])} testcase. "
      f"Trong {len(s['measured'])} bài có đủ số liệu oracle, macro TPR đạt {pct(s['macro_tpr'],2)}, macro TNR đạt {pct(s['macro_tnr'],2)}; micro TPR đạt {pct(s['micro_tpr'],2)} và micro TNR đạt {pct(s['micro_tnr'],2)}. "
      f"Có {s['both_08']}/{len(s['measured'])} bài đồng thời đạt TPR và TNR từ 0,8; {s['both_09']}/{len(s['measured'])} bài đồng thời đạt từ 0,9.")
    p(doc,
      "Kết quả cho thấy kiến trúc sinh lặp dựa trên feedback là khả thi trên cả CodeNet và Codeforces, nhưng chất lượng còn không đồng đều. "
      "Các rủi ro chính là nhãn lịch sử không tuyệt đối, reference consensus có thể cùng sai, checker không phù hợp có thể tạo false positive, và holdout nội bộ mới xuất hiện ở số ít run. "
      "Báo cáo vì vậy trình bày cả số liệu đạt được, giới hạn đo lường và kế hoạch kiểm chứng tiếp theo thay vì coi verdict lịch sử là ground truth tuyệt đối.")
    p(doc, "Từ khóa: sinh testcase tự động; mô hình ngôn ngữ lớn; mutation testing; test oracle; CodeNet; Codeforces; TPR; TNR.")

    doc.add_heading("ABSTRACT", level=1)
    p(doc,
      "This thesis studies an LLM-assisted pipeline that analyzes programming problem statements, synthesizes reference and faulty programs, generates tests from explicit coverage obligations, certifies expected outputs through multi-reference consensus, locks the resulting suite, and evaluates it against an external corpus of historically accepted and wrong submissions. The system targets both fidelity—the ability to preserve correct programs—and discriminability—the ability to reject faulty programs.")
    p(doc,
      f"The inspected workspace contains {s['problem_dirs']} problem directories and {num(s['manifest_unique'])} distinct source programs. The latest valid results include {s['ever_success_count']} locked suites with {num(s['cases_total'])} certified tests. Across {len(s['measured'])} problems with complete external-oracle measurements, macro TPR is {s['macro_tpr']:.4f} and macro TNR is {s['macro_tnr']:.4f}. The evidence supports the feasibility of feedback-guided generation while also showing uneven per-problem quality and the need for stronger label, checker, and oracle audits.")
    p(doc, "Keywords: automated test generation; large language models; mutation testing; test oracle; competitive programming.")

    doc.add_page_break()
    doc.add_heading("MỤC LỤC", level=1)
    add_toc(doc)
    doc.add_heading("DANH MỤC HÌNH", level=1)
    add_bullets(doc, [
        "Hình 2.1. Phân bố ngôn ngữ và trạng thái submission trong Project CodeNet.",
        "Hình 2.2. Khung sinh testcase lặp dựa trên feedback của CodeContests-O.",
        "Hình 2.3. Pareto TPR–TNR của các bộ dữ liệu CodeContests.",
        "Hình 2.4. Chuỗi chứng nhận trong mô hình coding agent làm test-suite auditor.",
        "Hình 3.1. Kiến trúc logic của Final Testcase Generator.",
        "Hình 6.1. Phân bố TPR–TNR theo từng bài.",
        "Hình 6.2. Macro TPR và TNR theo nguồn dữ liệu.",
        "Hình 6.3. Phân bố thời gian sinh và chứng nhận bộ test.",
        "Hình 6.4. Nguyên nhân thất bại trong các batch hoàn tất.",
    ])
    doc.add_heading("DANH MỤC TỪ VIẾT TẮT", level=1)
    add_table(doc, ["Từ viết tắt", "Tiếng Anh", "Diễn giải"], [
        ["AC", "Accepted", "Chương trình được nguồn dữ liệu ghi nhận là đúng"],
        ["WA", "Wrong Answer", "Chương trình cho kết quả sai"],
        ["CE", "Compilation Error", "Lỗi biên dịch"],
        ["RE", "Runtime Error", "Lỗi trong khi chạy"],
        ["TLE", "Time Limit Exceeded", "Vượt giới hạn thời gian"],
        ["LLM", "Large Language Model", "Mô hình ngôn ngữ lớn"],
        ["SUT", "System Under Test", "Đối tượng được kiểm thử"],
        ["TPR", "True Positive Rate", "Tỷ lệ chương trình đúng được giữ"],
        ["TNR", "True Negative Rate", "Tỷ lệ chương trình sai bị bắt"],
        ["SBST", "Search-Based Software Testing", "Kiểm thử phần mềm dựa trên tìm kiếm"],
    ], widths=[2.3, 5.2, 8.2], font_size=9)

    chapter(doc, "CHƯƠNG 1. GIỚI THIỆU")
    doc.add_heading("1.1 Bối cảnh", level=2)
    p(doc,
      "Kiểm thử phần mềm là hoạt động thực thi hoặc phân tích nhằm phát hiện sai khác giữa hành vi quan sát được và yêu cầu. Trong bài toán lập trình, một test gồm input, expected output và quy tắc so sánh. Một bộ test yếu có thể bỏ lọt chương trình sai; một bộ test hoặc checker sai có thể loại chương trình đúng. Sommerville xem kiểm thử là một phần của verification và validation, còn Copeland và Kaner nhấn mạnh thiết kế test theo miền dữ liệu, biên, lớp tương đương và lỗi có khả năng xảy ra [1]–[3].")
    p(doc,
      "Các nền tảng chấm bài sở hữu số lượng lớn đề và submission, nhưng hidden tests thường không công khai. Project CodeNet cung cấp hơn 14 triệu mẫu mã nguồn với metadata và nhãn verdict, tạo điều kiện dùng chương trình lịch sử làm tập đánh giá [8]. Dù vậy, nhãn Accepted chỉ nói rằng chương trình đã vượt bộ test tại thời điểm chấm; nghiên cứu audit gần đây chứng minh vẫn tồn tại accepted-but-buggy submissions [12]. Vì vậy, một pipeline sinh testcase cần vừa tận dụng nhãn lịch sử vừa công khai giới hạn của nhãn.")
    p(doc,
      "Đối với Codeforces, khóa luận sử dụng các submission từ bộ open-r1/codeforces-submissions trên Hugging Face. Dataset card công bố khoảng 12,6 triệu dòng, gồm source, problem_id, ngôn ngữ, verdict và thông tin tài nguyên, theo giấy phép CC BY 4.0 [13]. Riêng bài G1 — Yunli’s Subarray Queries (easy version), đề và các chương trình oracle được thu thập thủ công từ Codeforces; metadata lưu URL contest 2009, bài G1 [14].")
    doc.add_heading("1.2 Bài toán nghiên cứu", level=2)
    p(doc,
      "Với đề bài P, hệ thống cần xây dựng bộ test T sao cho mọi input hợp lệ theo ràng buộc, expected output đủ tin cậy, các trường hợp quan trọng được phủ và T phân biệt được nhiều chương trình sai mà không loại nhầm chương trình đúng. Bài toán có tính đa mục tiêu: tăng TNR thường dễ bằng cách thêm test khắt khe, nhưng expected output hoặc checker sai sẽ làm giảm TPR. Thời gian sinh cũng là ràng buộc thực tế.")
    add_note(doc,
      "Đối tượng của khóa luận là quy trình tạo và chứng nhận bộ test cho bài toán lập trình từ đề bài, mã nguồn ứng viên và hạ tầng thực thi. Báo cáo không đánh giá khả năng sinh lời giải cuối cho người học.",
      "Phạm vi")
    doc.add_heading("1.3 Mục tiêu", level=2)
    add_numbered(doc, [
        "Thiết kế kiến trúc module hóa cho LLM generator, judge, checker, oracle và cơ chế retry.",
        "Xây dựng chuỗi chứng nhận reference, mutant và testcase trước khi tính metric.",
        "Sinh lặp theo feedback từ chương trình sống sót và các nghĩa vụ phủ chưa đạt.",
        "Khóa test suite trước khi chấm oracle ngoài để ngăn rò rỉ dữ liệu đánh giá.",
        "Đánh giá định lượng bằng TPR, TNR, coverage, số testcase, tỷ lệ hoàn tất và thời gian sinh.",
        "Lưu artifact đủ chi tiết để truy vết từ số liệu tổng hợp về report và log gốc.",
    ])
    doc.add_heading("1.4 Câu hỏi nghiên cứu", level=2)
    add_table(doc, ["Mã", "Câu hỏi"], [
        ["RQ1", "Pipeline có tạo được bộ test hợp lệ, có khóa và có thể truy vết trên nhiều bài không?"],
        ["RQ2", "Bộ test giữ được bao nhiêu chương trình mang nhãn đúng và bắt được bao nhiêu chương trình mang nhãn sai?"],
        ["RQ3", "Chất lượng có khác nhau giữa CodeNet, Codeforces và từng bài không?"],
        ["RQ4", "Thời gian sinh có đáp ứng mốc mục tiêu 180 giây không?"],
        ["RQ5", "Những lỗi nào làm pipeline dừng trước khi có metric và đâu là rủi ro đối với độ tin cậy kết luận?"],
    ], widths=[1.8, 14.0], font_size=9.5)
    doc.add_heading("1.5 Đóng góp", level=2)
    add_bullets(doc, [
        "Một pipeline đầu cuối từ statement đến test suite được chứng nhận và đánh giá ngoài.",
        "Một test model có grammar, constraints, quan hệ, fault model và coverage obligations làm cầu nối giữa phân tích đề và sinh test.",
        "Một chuỗi chứng nhận nhiều tầng: validator, sample, reference consensus, witness và deterministic replay.",
        "Một giao thức đo tách development, holdout và external oracle, kèm test-suite lock.",
        f"Một tập bằng chứng hiện có gồm {s['ever_success_count']} bộ test khóa, {num(s['cases_total'])} testcase và kết quả oracle trên {len(s['measured'])} bài.",
    ])
    doc.add_heading("1.6 Cấu trúc báo cáo", level=2)
    p(doc,
      "Chương 2 trình bày thuật ngữ, nền tảng và công trình liên quan. Chương 3 mô tả phương pháp đề xuất. Chương 4 trình bày hiện thực codebase và dữ liệu. Chương 5 mô tả thiết kế thực nghiệm. Chương 6 báo cáo kết quả. Chương 7 thảo luận giới hạn, hướng phát triển và kết luận. Phụ lục cung cấp bảng theo từng bài và bản đồ artifact.")

    chapter(doc, "CHƯƠNG 2. CƠ SỞ LÝ THUYẾT, THUẬT NGỮ VÀ CÔNG TRÌNH LIÊN QUAN")
    doc.add_heading("2.1 Thuật ngữ nền tảng", level=2)
    terms = [
        ("Test case", "Một cặp input và expected output, kèm checker dùng để quyết định output thực tế có được chấp nhận hay không."),
        ("Test suite", "Tập testcase dùng chung để đánh giá một chương trình. Chất lượng suite phụ thuộc cả input, expected output và checker."),
        ("Test oracle", "Cơ chế xác định kết quả mong đợi hoặc verdict đúng. Trong dự án, oracle có hai nghĩa liên quan: reference consensus tạo expected output và tập submission lịch sử dùng đánh giá ngoài."),
        ("Validator", "Chương trình quyết định input có tuân theo grammar và constraints của statement. Validator không tính expected output."),
        ("Checker", "Hàm so sánh actual output và expected output. EXACT yêu cầu trùng tuyệt đối; TOKEN bỏ khác biệt khoảng trắng giữa token; FLOAT cần tolerance; PYTHON_SCRIPT xử lý semantics riêng."),
        ("Reference solution", "Chương trình được dùng để tạo output tham chiếu. Pipeline yêu cầu nhiều reference độc lập cùng đồng ý để giảm rủi ro từ một lời giải đơn."),
        ("Candidate AC", "Chương trình do LLM sinh với ý định đúng. Nhãn này tạm thời cho đến khi qua biên dịch, sample và các test được chứng nhận."),
        ("Candidate WA hoặc mutant", "Chương trình cố ý sai theo một fault model. Mutant chỉ có giá trị đo lường sau khi tồn tại witness hợp lệ chứng minh sai."),
        ("Witness", "Input hợp lệ làm lộ sai khác giữa mutant và reference consensus. Witness phải tái hiện ổn định qua các lần chạy."),
        ("Quarantine", "Vùng lưu test bị loại vì input không hợp lệ, expected output bất đồng, sample không khớp hoặc lý do chứng nhận khác."),
        ("Coverage obligation", "Mục tiêu kiểm thử có tên, chẳng hạn min, max, duplicate, overflow, vô nghiệm hoặc cấu trúc đặc biệt. Coverage trong dự án là coverage của obligation, không phải line coverage của submission."),
        ("Development set", "Nhóm candidate được chấm trong vòng lặp và có thể tạo feedback cho LLM."),
        ("Holdout set", "Nhóm mutant được giữ khỏi feedback và chỉ chấm sau khi khóa suite, dùng kiểm tra khả năng tổng quát trong pool nội bộ."),
        ("External oracle", "Tập chương trình Accepted và Wrong Answer lấy từ CodeNet hoặc Codeforces, chỉ được đánh giá sau khi suite khóa."),
        ("Test-suite lock", "Artifact lưu danh sách và hash suite cuối. Khóa giúp phát hiện việc thay đổi test sau khi nhìn kết quả oracle."),
        ("TPR", "Số chương trình mang nhãn đúng PASS chia cho số chương trình đúng đánh giá được. Chỉ số đo fidelity của suite."),
        ("TNR", "Số chương trình mang nhãn sai bị WA, RE hoặc TLE chia cho số chương trình sai đánh giá được. Chỉ số đo discriminability."),
        ("Macro average", "Tính metric cho từng bài rồi lấy trung bình; mỗi bài có trọng số ngang nhau."),
        ("Micro average", "Cộng tử số và mẫu số qua mọi bài rồi chia; bài có nhiều chương trình đánh giá được có trọng số cao hơn."),
        ("False positive trong ngữ cảnh này", "Một chương trình mang nhãn đúng bị suite loại. Nguyên nhân có thể là test sai, checker sai, toolchain khác hoặc chính nhãn Accepted không còn đúng."),
        ("Survivor", "Mutant hoặc chương trình Wrong Answer vẫn PASS; đây là tín hiệu để tạo feedback và bổ sung test nhắm vào lớp lỗi chưa bị bắt."),
    ]
    for term, explanation in terms:
        para = doc.add_paragraph()
        para.paragraph_format.keep_together = True
        para.add_run(term + ". ").bold = True
        para.add_run(explanation)

    doc.add_heading("2.2 Thiết kế testcase truyền thống", level=2)
    p(doc,
      "Phân vùng tương đương chia miền input thành các lớp mà hệ thống được kỳ vọng xử lý giống nhau; phân tích giá trị biên tập trung vào điểm đầu, điểm cuối và vùng ngay quanh biên. Decision table phù hợp khi output phụ thuộc tổ hợp điều kiện; state transition phù hợp khi hành vi phụ thuộc trạng thái. Các kỹ thuật này tạo nền cho test model của dự án: LLM phải mô tả miền, quan hệ và lớp lỗi trước khi đề xuất input [2], [3].")
    p(doc,
      "Search-based software testing biểu diễn mục tiêu kiểm thử thành hàm thích nghi rồi tìm input tối ưu. McMinn tổng hợp các hướng sử dụng thuật toán metaheuristic cho sinh dữ liệu kiểm thử và nêu rõ khó khăn ở biểu diễn, landscape và oracle [4]. Dự án hiện tại không tối ưu trực tiếp branch distance; nó dùng LLM để đề xuất cấu trúc input và dùng feedback thực thi để chuyển hướng vòng sau.")
    doc.add_heading("2.3 Mutation testing", level=2)
    p(doc,
      "Mutation testing tạo các phiên bản sai của chương trình và đo khả năng bộ test tiêu diệt chúng. Jia và Harman chỉ ra mutation score là tín hiệu hữu ích nhưng chịu ảnh hưởng của equivalent mutants và chi phí chạy lớn [5]. Trong dự án, mutant không được coi là sai chỉ vì prompt yêu cầu sinh sai: hệ thống đòi witness, validator, reference consensus và hai lần replay. Quy tắc này trực tiếp xử lý rủi ro equivalent hoặc non-reproducible mutant.")
    doc.add_heading("2.4 LLM và vòng lặp feedback", level=2)
    p(doc,
      "CodaMosa kết hợp search-based generation với Codex khi coverage bị chững, cho thấy LLM có thể cung cấp seed giúp vượt coverage plateau [6]. CoverUp đưa code context, coverage và lỗi thực thi trở lại prompt qua nhiều vòng và báo cáo cải thiện coverage trên Python modules [7]. Hai công trình này củng cố lựa chọn kiến trúc lặp: LLM sinh đề xuất, hệ thống chạy thật, rồi feedback phải chứa survivor, lỗi và mục tiêu chưa phủ.")
    doc.add_heading("2.5 Dữ liệu chương trình và verdict lịch sử", level=2)
    p(doc,
      "Project CodeNet có 13.916.868 submission cho 4.053 bài và 55 ngôn ngữ; 53,6% submission mang nhãn Accepted và 29,5% mang nhãn Wrong Answer [8]. Quy mô và metadata làm CodeNet phù hợp để lấy cả positive lẫn negative programs. Đối với Codeforces, nguồn submission chính là open-r1/codeforces-submissions [13]; G1 là trường hợp được crawl thủ công từ trang bài chính thức [14]. Phân bố ngôn ngữ và trạng thái không đồng đều, nên kết quả theo bài và theo ngôn ngữ cần được đọc cùng mẫu số.")
    add_picture(doc, ASSET_DIR / "codenet-distribution.png", 5.8, "2.1",
                "Phân bố submission theo ngôn ngữ và trạng thái trong Project CodeNet",
                "Puri và cộng sự [8, Fig. 1]; trích nguyên hình, chỉ cắt lề trang")
    doc.add_heading("2.6 Các bộ test cho competitive programming", level=2)
    p(doc,
      "CodeContests được giới thiệu cùng AlphaCode như một tập bài và lời giải cho nghiên cứu sinh mã [9]. CodeContests+ dùng agent LLM tạo test và đánh giá trên 1,72 triệu submission mang nhãn pass/fail, nhấn mạnh rằng chất lượng test ảnh hưởng trực tiếp đến độ chính xác đánh giá [10]. CodeContests-O mở rộng theo hướng feedback-driven: chạy test trên tập lời giải đúng/sai đã biết, phân tích false positive và false negative, rồi lặp lại quá trình sinh [11].")
    add_picture(doc, ASSET_DIR / "codecontests-o-framework.png", 6.25, "2.2",
                "Khung sinh testcase lặp dựa trên feedback của CodeContests-O",
                "Cai và cộng sự [11, Fig. 2]; trích nguyên hình, chỉ cắt lề trang")
    p(doc,
      "CodeContests-O định nghĩa TPR là tỷ lệ lời giải đúng được chấp nhận và TNR là tỷ lệ lời giải sai bị loại, tương thích với cách đọc metric trong khóa luận. Trên toàn bộ khoảng 11 triệu lời giải, bài báo báo cáo TPR 89,37% và TNR 90,89% sau ba vòng refinement [11]. Các con số này là mốc tham khảo khái niệm; chúng không phải baseline trực tiếp vì khác bộ dữ liệu, số bài, nguồn chương trình, checker, ngân sách và quy trình chứng nhận.")
    add_picture(doc, ASSET_DIR / "codecontests-o-pareto.png", 4.2, "2.3",
                "Pareto TPR–TNR của CodeContests, HardTests, CodeContests+ và CodeContests-O",
                "Cai và cộng sự [11, Fig. 1]; trích nguyên hình, chỉ cắt lề trang")
    doc.add_heading("2.7 Audit bộ test và độ tin cậy của nhãn", level=2)
    p(doc,
      "Xie và cộng sự xem coding agents như test-suite auditors. Chuỗi chứng nhận của họ gồm consensus oracle từ nhiều accepted solutions độc lập, brute-force adjudication khi bất đồng và validator nghiêm ngặt cho tính hợp lệ input [12]. Nghiên cứu phát hiện 589 accepted-but-buggy submissions trong 20.375 submission AtCoder ở một agent arm và ít nhất 906 ở hợp năm agent. Kết quả này là bằng chứng trực tiếp rằng Accepted không nên được xem là ground truth tuyệt đối.")
    add_picture(doc, ASSET_DIR / "agentic-certification-chain.png", 6.25, "2.4",
                "Mô hình coding agent làm test-suite auditor và chuỗi chứng nhận",
                "Xie và cộng sự [12, Fig. 1]; trích nguyên hình, chỉ cắt lề trang")
    doc.add_heading("2.8 Nguồn tham khảo đã sử dụng và vai trò trong đề tài", level=2)
    add_table(doc, ["Nguồn", "Nội dung sử dụng", "Cách áp dụng trong khóa luận"], [
        ["[1] Sommerville", "Verification, validation, testing process", "Đặt bối cảnh kỹ nghệ phần mềm và mục tiêu kiểm thử"],
        ["[2] Copeland", "Equivalence, boundary, decision table, state transition", "Cấu trúc coverage obligation và phân tích miền"],
        ["[3] Kaner et al.", "Thực hành thiết kế test và oracle", "Giải thích test case, checker và rủi ro loại nhầm"],
        ["[4] McMinn", "Search-based test-data generation", "Đối chiếu cơ chế tìm input bằng feedback"],
        ["[5] Jia–Harman", "Mutation testing và equivalent mutant", "Thiết kế mutant certification và cách đọc TNR"],
        ["[6] CodaMosa", "LLM hỗ trợ khi search bị chững", "Lý do dùng LLM kết hợp feedback thực thi"],
        ["[7] CoverUp", "Coverage-guided iterative prompting", "Lý do đưa mục tiêu chưa phủ và lỗi về prompt"],
        ["[8] CodeNet", "Quy mô, metadata, phân bố verdict", "Nguồn chương trình lịch sử và hình 2.1"],
        ["[9] CodeContests", "Bộ dữ liệu competitive programming", "Bối cảnh cho các bộ test tổng hợp"],
        ["[10] CodeContests+", "Agent sinh test, đánh giá 1,72 triệu submission", "Mốc liên quan cho test generation quy mô lớn"],
        ["[11] CodeContests-O", "Feedback loop, TPR/TNR, Pareto", "Đối chiếu gần nhất với phương pháp và hình 2.2–2.3"],
        ["[12] Test-suite auditors", "Consensus, validator, accepted-but-buggy", "Cơ sở cho chuỗi chứng nhận và cảnh báo nhãn; hình 2.4"],
        ["[13] open-r1 Codeforces", "12,6 triệu submission, source và verdict", "Nguồn submission oracle cho các bài Codeforces cf*"],
        ["[14] Codeforces G1", "Đề G1, contest 2009", "Nguồn bài và oracle được crawl thủ công"],
    ], widths=[3.1, 5.8, 6.9], font_size=8.4)
    # add_table leaves a spacer paragraph. Removing it here prevents an empty
    # page when the source-mapping table exactly fills the preceding page.
    spacer = doc.paragraphs[-1]
    if not spacer.text:
        spacer._element.getparent().remove(spacer._element)

    chapter(doc, "CHƯƠNG 3. PHƯƠNG PHÁP ĐỀ XUẤT")
    doc.add_heading("3.1 Mô hình bài toán", level=2)
    p(doc,
      "Gọi P là đề bài; V(x) là validator; C(y, ŷ) là checker; R={r1,…,rk} là tập reference đã chứng nhận; M là tập mutant; T là test suite; O+ và O− lần lượt là các chương trình lịch sử mang nhãn Accepted và Wrong Answer. Một testcase (x,y) được nhận khi V(x)=true và các reference đánh giá được đồng thuận về y theo C. Suite T được tối ưu theo ba nhóm mục tiêu: hợp lệ và tin cậy, coverage obligations, và khả năng phân biệt chương trình.")
    add_note(doc,
      "TPR(T)=|{p∈O+: p PASS T}|/|O+ đánh giá được|; TNR(T)=|{p∈O−: p bị WA/RE/TLE}|/|O− đánh giá được|. CE và lỗi hạ tầng không được tính vào mẫu số logic.",
      "Hàm đánh giá")
    doc.add_heading("3.2 Kiến trúc tổng thể", level=2)
    # Architecture table acts as a stable, editable native diagram.
    arch = add_table(doc, ["Tầng", "Thành phần", "Đầu ra chính"], [
        ["1. Nạp bài", "statement, metadata, checker, runtime", "ProblemSpec"],
        ["2. Phân tích", "LLM tạo grammar, constraints, relations, fault model", "test-model.json"],
        ["3. Sinh chương trình", "candidate AC/WA và validator", "generated/*"],
        ["4. Chứng nhận", "compile, sample, witness, replay, consensus", "reference/mutant certification"],
        ["5. Sinh test", "baseline, repair, deduplicate, expected output", "certified cases"],
        ["6. Tinh chỉnh", "survivor + obligation chưa phủ → feedback", "iteration snapshots"],
        ["7. Khóa và đánh giá", "suite lock, holdout, external oracle", "report.json và execution logs"],
    ], widths=[3.1, 7.7, 5.0], font_size=9)
    add_caption(doc, "3.1", "Kiến trúc logic của Final Testcase Generator",
                "tác giả tổng hợp từ main.py và testgen/pipeline.py")
    doc.add_heading("3.3 Test model", level=2)
    p(doc,
      "Test model là biểu diễn trung gian có cấu trúc. Grammar mô tả token và dòng; constraints ghi miền và quan hệ; equivalence classes gom các hành vi tương tự; boundary cases liệt kê giá trị biên; metamorphic relations mô tả quan hệ giữa nhiều input/output; fault model chỉ ra lỗi cài đặt có khả năng xảy ra; coverage obligations chuyển các phân tích trên thành mục tiêu có mã và độ ưu tiên. Nhờ đó prompt sinh test không bắt đầu từ một yêu cầu chung chung mà từ danh sách mục tiêu có thể đo.")
    doc.add_heading("3.4 Chứng nhận reference", level=2)
    add_numbered(doc, [
        "Sinh nhiều candidate AC ở các ngôn ngữ được hỗ trợ.",
        "Biên dịch trong runtime tương ứng; loại CE.",
        "Chạy official samples và so theo checker; loại candidate không khớp.",
        "Giữ tập reference đủ kích thước tối thiểu; nếu không đủ, pipeline dừng trước khi tạo metric.",
        "Khi gán expected output cho test mới, yêu cầu ít nhất hai reference đánh giá được đồng ý theo checker.",
    ])
    doc.add_heading("3.5 Chứng nhận mutant", level=2)
    p(doc,
      "Mutant được sinh theo fault model hoặc biến đổi từ reference. Mỗi mutant phải có witness được validator chấp nhận. Reference consensus cung cấp expected output; mutant phải tạo verdict sai. Cặp mutant–witness được chạy lặp lại hai lần để loại bất định. Pipeline cố gắng thay thế mutant không đạt qua tối đa bốn vòng và chỉ tiếp tục khi có ít nhất ba mutant được chứng nhận.")
    doc.add_heading("3.6 Sinh và tinh chỉnh testcase", level=2)
    add_numbered(doc, [
        "Tạo baseline từ sample, biên và coverage obligation.",
        "Sửa input sai định dạng theo recipe; validator vẫn là cổng quyết định cuối.",
        "Chuẩn hóa và loại trùng input.",
        "Tính expected output bằng sample hoặc reference consensus.",
        "Chấm development candidates, ghi TPR/TNR và danh sách survivor.",
        "Sinh feedback gồm false positive, false negative, execution error và nghĩa vụ chưa phủ.",
        "Tạo kế hoạch vòng tiếp theo nhắm survivor hoặc gap coverage.",
        "Dừng khi đạt tiêu chí chất lượng hoặc hết số vòng cấu hình.",
    ])
    doc.add_heading("3.7 Khóa suite và chống rò rỉ", level=2)
    p(doc,
      "Sau vòng refinement, hệ thống ghi test-suite.lock.json chứa hash của suite. Holdout và oracle ngoài chỉ chạy sau mốc này. Pipeline kiểm tra khóa trước và sau chấm để bảo đảm test không thay đổi theo kết quả oracle. Đây là ranh giới quan trọng giữa dữ liệu dùng phát triển và dữ liệu dùng báo cáo.")
    doc.add_heading("3.8 Cơ chế judge và checker", level=2)
    p(doc,
      "Judge abstraction hỗ trợ Docker local, remote judge qua SSH và native smoke test. Docker đặt giới hạn bộ nhớ, CPU, process và timeout; retrying judge tách lỗi hạ tầng khỏi verdict. Checker abstraction hỗ trợ so sánh exact, token, decimal, float, không phân biệt hoa thường và script Python. Việc tách checker khỏi judge cho phép cùng execution output được đánh giá theo semantics của từng bài.")

    chapter(doc, "CHƯƠNG 4. HIỆN THỰC HỆ THỐNG VÀ DỮ LIỆU")
    doc.add_heading("4.1 Cấu trúc codebase", level=2)
    add_table(doc, ["Module", "Trách nhiệm"], [
        ["main.py", "CLI, batch runner, composition root và cấu hình dependency"],
        ["testgen/domain.py", "Dataclass, protocol và kiểu dữ liệu dùng xuyên pipeline"],
        ["testgen/pipeline.py", "Orchestrator của toàn bộ quy trình sinh–chứng nhận–đánh giá"],
        ["testgen/ollama.py", "Backend Ollama, schema prompt và xử lý phản hồi"],
        ["testgen/claude.py", "Backend Claude Code"],
        ["testgen/chatgpt.py", "Backend ChatGPT"],
        ["testgen/fallback.py", "Fallback theo loại nhiệm vụ"],
        ["testgen/checker.py; testgen/checkers/*", "Resolver và các chiến lược checker"],
        ["testgen/docker_judge.py", "Biên dịch/chạy cô lập bằng Docker"],
        ["testgen/remote_judge.py", "Thực thi trên KHOI qua SSH"],
        ["testgen/native_judge.py", "Smoke test cục bộ"],
        ["testgen/retrying_judge.py", "Retry lỗi hạ tầng có kiểm soát"],
        ["testgen/codenet.py", "Đọc metadata và oracle CodeNet"],
        ["testgen/coverage.py", "Tính coverage obligation"],
        ["testgen/progress.py", "Ghi tiến độ và artifact"],
        ["testgen/runtime.py", "Ánh xạ ngôn ngữ sang compile/run command"],
        ["testgen/input_recipe.py", "Sửa và chuẩn hóa input được sinh"],
        ["testgen/source_edits.py", "Áp dụng chỉnh sửa source có kiểm soát"],
    ], widths=[6.2, 9.6], font_size=8.8)
    doc.add_heading("4.2 Cấu hình thực nghiệm hiện tại", level=2)
    add_table(doc, ["Tham số", "Giá trị", "Vai trò"], [
        ["Candidate AC", "3", "Reference pool ban đầu"],
        ["Candidate WA", "5", "Mutant pool ban đầu"],
        ["Test/vòng", "10", "Ngân sách sinh testcase mỗi vòng"],
        ["Số vòng tối đa", "2", "Giới hạn refinement hiện tại"],
        ["Test tối thiểu", "18", "Điều kiện dừng sớm"],
        ["Critical coverage", "1,0", "Phải phủ toàn bộ nghĩa vụ critical"],
        ["Normal coverage", "0,9", "Mục tiêu nghĩa vụ normal"],
        ["Holdout ratio", "0,2", "Tỷ lệ mutant nội bộ giữ kín khi đủ pool"],
        ["Mutant tối thiểu", "3", "Cổng tiếp tục pipeline"],
        ["Replacement rounds", "4", "Số vòng thay mutant không đạt"],
        ["Run workers", "8", "Song song thực thi"],
        ["Compile workers", "2", "Song song biên dịch"],
        ["Run timeout", "5 giây", "Giới hạn một execution"],
        ["Compile timeout", "90 giây", "Giới hạn biên dịch"],
        ["Sandbox", "512 MiB; 1 CPU; 64 PID", "Giới hạn tài nguyên"],
    ], widths=[4.4, 3.4, 8.0], font_size=8.8)
    doc.add_heading("4.3 Dữ liệu bài toán", level=2)
    datasets = Counter(row["dataset"] for row in e["data_rows"])
    add_table(doc, ["Hạng mục", "Số lượng", "Ghi chú"], [
        ["Thư mục bài toán", s["problem_dirs"], f"{datasets.get('CodeNet',0)} CodeNet; {datasets.get('Codeforces',0)} Codeforces"],
        ["statement.md", s["statements"], "Đầu vào cho phân tích đề"],
        ["metadata.json", s["metadata_count"], "Thông tin nguồn và bài"],
        ["checker.json", s["checker_count"], "Cấu hình checker riêng"],
        ["Oracle bundle", s["oracle_bundle_count"], "Có manifest và source"],
        ["Mục manifest", num(s["manifest_entries"]), "Có thể gồm lặp theo bundle"],
        ["Source duy nhất theo SHA-256", num(s["manifest_unique"]), "Đơn vị đếm không trùng"],
        ["Accepted", num(s["manifest_accepted"]), "Nhãn lịch sử"],
        ["Wrong Answer", num(s["manifest_wrong"]), "Nhãn lịch sử"],
    ], widths=[5.2, 3.1, 7.5], font_size=9)
    p(doc,
      "Trong 32 bài Codeforces, 31 thư mục có mã cf* sử dụng submission oracle từ open-r1/codeforces-submissions; trường dataset trong manifest lưu đúng định danh này [13]. Bài G1 cũng thuộc Codeforces nhưng được thu thập thủ công: metadata trỏ đến contest 2009/problem/G1 và script fetch_cf_g1.py lưu lại quy trình lấy statement [14].")
    doc.add_heading("4.4 Checker profile", level=2)
    p(doc,
      "Có 42 bài khai báo checker profile: 22 TOKEN, 17 EXACT và 3 PYTHON_SCRIPT; 30 bài không có profile riêng nên dùng mặc định TOKEN. Sự phân bố này ảnh hưởng trực tiếp đến metric. TOKEN phù hợp với output dạng token nhưng không thể thay thế special checker khi tồn tại nhiều output hợp lệ; PYTHON_SCRIPT được dùng cho semantics riêng.")
    doc.add_heading("4.5 Artifact và khả năng truy vết", level=2)
    p(doc,
      "Một run thành công lưu statement, run-config, test-model, code sinh, chứng nhận reference/mutant, case plan, feedback, execution logs, testcase .in/.out, report.json và test-suite.lock.json. Batch lưu batch-report.json và log từng bài. Cấu trúc này cho phép đi từ một điểm bất thường trong bảng tổng hợp về testcase, source và verdict cụ thể.")
    chapter(doc, "CHƯƠNG 5. THIẾT KẾ THỰC NGHIỆM")
    doc.add_heading("5.1 Đơn vị phân tích", level=2)
    p(doc,
      "Đơn vị chính là một bài có report pipeline đúng schema và test-suite.lock.json. Khi một bài có nhiều run hợp lệ, báo cáo dùng run thành công gần nhất. Trạng thái lần thử gần nhất vẫn được lưu riêng để phản ánh lỗi mới xảy ra. Batch chỉ được dùng khi finished_at tồn tại và completed bằng total; batch dừng giữa chừng không đóng góp số liệu.")
    doc.add_heading("5.2 Tập đánh giá", level=2)
    p(doc,
      "Development candidates được dùng trong refinement. Holdout mutants chỉ chấm sau lock và hiện có ở 5 run. External oracle gồm chương trình lịch sử từ CodeNet hoặc Codeforces. Tập này lớn hơn candidate pool nội bộ và cung cấp phép đo ngoài sau khi suite đã cố định. G1 được xếp đúng vào nhóm Codeforces; khác biệt duy nhất là dữ liệu của bài này được crawl thủ công thay vì lấy từ bộ open-r1 [13], [14].")
    doc.add_heading("5.3 Chỉ số", level=2)
    add_table(doc, ["Chỉ số", "Cách tính", "Ý nghĩa"], [
        ["TPR", "AC PASS / AC đánh giá được", "Fidelity; càng cao càng ít loại nhầm theo nhãn"],
        ["TNR", "WA bị WA/RE/TLE / WA đánh giá được", "Discriminability; càng cao càng bắt nhiều sai"],
        ["Coverage", "Obligation đã phủ / tổng obligation", "Mức hoàn tất test model"],
        ["Số testcase", "Case được chứng nhận trong suite khóa", "Quy mô suite"],
        ["Latency", "Từ startup đến suite chứng nhận và holdout", "Không gồm optional external oracle"],
        ["Completion", "Run/batch có đủ artifact kết thúc", "Độ ổn định vận hành"],
    ], widths=[2.7, 6.1, 7.0], font_size=9)
    doc.add_heading("5.4 Giả thuyết và tiêu chí đọc", level=2)
    add_bullets(doc, [
        "H1: macro TPR và macro TNR đều đạt mức hữu dụng khoảng 0,8 trở lên.",
        "H2: ít nhất hai phần ba bài đo đủ đồng thời đạt TPR và TNR từ 0,8.",
        "H3: đa số run có generation latency không vượt 180 giây.",
        "H4: feedback loop và certification tạo được artifact hoàn chỉnh trên phần lớn lượt trong batch hoàn tất.",
    ])
    doc.add_heading("5.5 Đối chiếu với nghiên cứu liên quan", level=2)
    p(doc,
      "CodeContests-O là đối chiếu gần nhất vì cùng dùng TPR/TNR và feedback từ lời giải đúng/sai [11]. Tuy nhiên, so sánh chỉ mang tính bối cảnh: CodeContests-O chạy trên 1,1×10^7 lời giải và dataset riêng, trong khi thí nghiệm hiện tại có 59 bài đo đủ và 2.235 execution logic trong micro metric. Không thực hiện kiểm định thống kê giữa hai hệ thống vì đơn vị mẫu và protocol không tương đương.")
    doc.add_heading("5.6 Kiểm soát sai lệch", level=2)
    add_bullets(doc, [
        "Loại batch chưa hoàn tất để tránh trộn snapshot bị dừng giữa chừng.",
        "Yêu cầu suite lock để tránh sửa test sau khi nhìn oracle.",
        "Loại CE và lỗi hạ tầng khỏi mẫu số logic thay vì gán thành WA.",
        "Báo cáo cả macro và micro để tránh một cách gộp duy nhất che khuất phân bố.",
        "Tách số liệu theo nguồn và liệt kê từng bài ở phụ lục.",
        "Không xem nhãn lịch sử là chân lý tuyệt đối; mọi kết luận đều dùng cụm từ ‘theo nhãn’."
    ])

    chapter(doc, "CHƯƠNG 6. KẾT QUẢ VÀ THẢO LUẬN")
    doc.add_heading("6.1 Mức độ hoàn tất", level=2)
    add_table(doc, ["Kết quả", "Giá trị", "Diễn giải"], [
        ["Batch hoàn tất", len(s["complete_batches"]), "Được đưa vào thống kê"],
        ["Batch bị loại", len(s["excluded_batches"]), "Thiếu finished_at hoặc completed < total"],
        ["Lượt trong batch hoàn tất", s["complete_batch_attempts"], f"{s['complete_batch_success']} thành công; {s['complete_batch_failed']} thất bại"],
        ["Bài có suite khóa", s["ever_success_count"], "Run thành công gần nhất theo bài"],
        ["Bài có đủ TPR/TNR", len(s["measured"]), "Đưa vào tổng hợp chất lượng"],
        ["Testcase chứng nhận", num(s["cases_total"]), f"Trung bình {num(s['cases_avg'],2)} test/bài"],
        ["Coverage 100%", f"{s['coverage_full']}/{s['ever_success_count']}", "Theo obligation trong test model"],
        ["Dừng vì đủ chất lượng", f"{s['quality_stop']}/{s['ever_success_count']}", "Đạt điều kiện dừng sớm"],
        ["Dừng do hết vòng", f"{s['iteration_stop']}/{s['ever_success_count']}", "Còn ít nhất một điều kiện chưa đạt"],
    ], widths=[5.0, 3.0, 7.8], font_size=9)
    p(doc,
      "RQ1 được trả lời tích cực ở mức prototype: pipeline tạo được suite khóa cho 60 bài và ghi đủ artifact để truy vết. Tuy nhiên, 17/81 lượt trong các batch hoàn tất vẫn dừng trước metric, nên khả năng vận hành chưa ổn định tuyệt đối. Việc 33/60 run thành công dừng do hết hai vòng cũng cho thấy ngân sách refinement hiện tại thường quá ngắn.")
    doc.add_heading("6.2 Chất lượng oracle tổng thể", level=2)
    add_table(doc, ["Metric", "Macro", "Micro", "Tử số/mẫu số micro"], [
        ["TPR", pct(s["macro_tpr"],2), pct(s["micro_tpr"],2), f"{s['pos_pass']}/{s['pos_eval']}"],
        ["TNR", pct(s["macro_tnr"],2), pct(s["micro_tnr"],2), f"{s['neg_killed']}/{s['neg_eval']}"],
    ], widths=[3.2, 3.2, 3.2, 6.2], font_size=10)
    p(doc,
      f"Macro TPR {pct(s['macro_tpr'],2)} và macro TNR {pct(s['macro_tnr'],2)} đều vượt 0,8, hỗ trợ H1. Có {s['both_08']}/{len(s['measured'])} bài, tương đương {pct(s['both_08']/len(s['measured']),1)}, đạt đồng thời hai metric từ 0,8; tỷ lệ này nhỉnh hơn hai phần ba và hỗ trợ H2. Nhóm rất mạnh gồm {s['both_09']}/{len(s['measured'])} bài đạt đồng thời từ 0,9.")
    add_picture(doc, charts["scatter"], 6.35, "6.1",
                "Phân bố TPR–TNR của từng bài có đủ oracle metric",
                "tác giả tổng hợp từ report.json; đường 0,8 và 0,9 là ngưỡng đọc kết quả")
    p(doc,
      "Phân bố cho thấy trade-off theo bài: một số suite bắt gần hết WA nhưng làm giảm TPR, trong khi một số giữ toàn bộ AC nhưng bỏ lọt nhiều WA. Vì vậy, chỉ báo trung bình không thay thế việc điều tra các điểm nằm ngoài vùng trên-phải.")
    doc.add_heading("6.3 Kết quả theo nguồn dữ liệu", level=2)
    group_rows = []
    for group in ("CodeNet", "Codeforces", "Khác"):
        runs = [r for r in s["runs"] if dataset_of(r["problem"]) == group]
        measured = [r for r in runs if r["oracle_tpr"] is not None and r["oracle_tnr"] is not None]
        if not runs:
            continue
        group_rows.append([
            group, len(runs), len(measured),
            pct(statistics.mean(r["oracle_tpr"] for r in measured),2) if measured else "—",
            pct(statistics.mean(r["oracle_tnr"] for r in measured),2) if measured else "—",
            f"{sum(r['oracle_tpr'] >= .8 and r['oracle_tnr'] >= .8 for r in measured)}/{len(measured)}" if measured else "—",
            f"{sum(r['coverage'] == 1 for r in runs)}/{len(runs)}",
        ])
    add_table(doc, ["Nguồn", "Run", "Đo đủ", "TPR TB", "TNR TB", "Cả hai ≥0,8", "Coverage 100%"],
              group_rows, widths=[2.7,1.6,1.8,2.2,2.2,2.7,2.8], font_size=8.7)
    add_picture(doc, charts["performance"], 6.25, "6.2",
                "Macro TPR và TNR theo nguồn dữ liệu",
                "tác giả tổng hợp từ report.json")
    p(doc,
      "Sau khi gộp đúng G1 vào Codeforces, hai nguồn vẫn có macro TPR gần nhau; TNR CodeNet cao hơn Codeforces. Chênh lệch chưa đủ để kết luận dataset nào dễ hơn vì mỗi nhóm có tập bài, checker và oracle khác nhau. G1 là một bài Codeforces được crawl thủ công, còn các bài cf* dùng submission từ open-r1/codeforces-submissions [13], [14].")
    doc.add_heading("6.4 Coverage và quy mô suite", level=2)
    p(doc,
      f"Tổng cộng {num(s['cases_total'])} testcase được chứng nhận, trung bình {num(s['cases_avg'],2)} test mỗi bài. {s['coverage_full']}/{s['ever_success_count']} bài phủ toàn bộ obligation. Coverage 100% không bảo đảm TNR cao vì obligation được suy ra từ test model: nếu fault model thiếu một lớp lỗi, suite có thể hoàn thành mọi mục tiêu đã biết nhưng vẫn bỏ lọt WA lịch sử. Ngược lại, coverage thấp không nhất thiết làm TPR/TNR thấp nếu các test ít nhưng trúng lớp lỗi chính.")
    doc.add_heading("6.5 Thời gian sinh", level=2)
    p(doc,
      f"Có {len(s['generation'])} run ghi latency. Trung vị là {num(s['generation_median'])} giây, trung bình {num(s['generation_avg'])} giây; {s['generation_le_180']}/{len(s['generation'])} run, tương đương {pct(s['generation_le_180']/len(s['generation']),1)}, không vượt 180 giây. Như vậy đa số run đạt mốc và hỗ trợ H3, nhưng phần đuôi dài kéo trung bình lên. Latency không bao gồm optional external oracle.")
    add_picture(doc, charts["latency"], 6.35, "6.3",
                "Phân bố thời gian sinh và chứng nhận bộ test; đường đỏ nằm đúng tại ranh giới 180 giây",
                "tác giả tổng hợp từ generation_latency trong report.json")
    slow = sorted([r for r in s["runs"] if r["generation_seconds"] is not None], key=lambda r:r["generation_seconds"], reverse=True)[:10]
    add_table(doc, ["Bài", "Giây", "Test", "Coverage", "Lý do dừng"], [
        [r["problem"], num(r["generation_seconds"]), r["cases"], pct(r["coverage"]), "Đủ chất lượng" if r["stop_reason"] == "all_quality_criteria_met" else "Hết vòng"] for r in slow
    ], widths=[3.0,2.2,2.0,2.7,5.9], font_size=8.8)
    doc.add_heading("6.6 Các bài cần ưu tiên điều tra", level=2)
    weak = sorted(s["measured"], key=lambda r:(min(r["oracle_tpr"],r["oracle_tnr"]),(r["oracle_tpr"]+r["oracle_tnr"])/2))[:14]
    add_table(doc, ["Bài", "Nguồn", "TPR", "TNR", "Coverage", "Test", "Ưu tiên"], [
        [r["problem"], dataset_of(r["problem"]), pct(r["oracle_tpr"]), pct(r["oracle_tnr"]), pct(r["coverage"]), r["cases"],
         "Bảo toàn AC" if r["oracle_tpr"] < r["oracle_tnr"] else "Bắt WA"] for r in weak
    ], widths=[2.2,2.2,1.8,1.8,2.3,1.4,4.1], font_size=8.2)
    p(doc,
      "Các bài TPR thấp phải được phân tích trước theo checker, expected output, runtime và từng AC bị loại; không nên mặc định testcase quá khó. Các bài TNR thấp cần phân cụm survivor theo lỗi logic rồi bổ sung obligation nhắm mục tiêu. p03044 là trường hợp đặc biệt: coverage model bằng 0 và TNR rất thấp, trong khi TPR bằng 1; đây là dấu hiệu suite quá nhỏ hoặc test model không tạo obligation hữu ích.")
    doc.add_heading("6.7 Lỗi vận hành", level=2)
    failures = [a for a in e["complete_attempts"] if a["status"] == "failed"]
    fc = Counter(a["failure_reason"] for a in failures)
    add_table(doc, ["Nhóm lỗi", "Số lượt", "Hàm ý"], [
        [reason, count, {
            "Không sinh đủ AC hoặc validator":"Dừng trước khi thiết lập đầu vào tin cậy",
            "Không đủ mutant WA được chứng nhận":"Witness không tái hiện hoặc pool sai quá yếu",
            "Không đủ reference qua biên dịch và sample":"Không tạo được consensus an toàn",
            "Remote judge không sẵn sàng":"Sự cố hạ tầng; không phải verdict logic",
        }.get(reason,"Cần đọc log chi tiết")] for reason,count in fc.most_common()
    ], widths=[6.3,2.0,7.5], font_size=9)
    add_picture(doc, charts["failures"], 6.25, "6.4",
                "Nguyên nhân thất bại trong các batch hoàn tất",
                "tác giả phân loại từ log của từng lượt")
    p(doc,
      "Bảy lượt thất bại do không sinh đủ AC hoặc validator, sáu lượt do không đủ mutant được chứng nhận, ba lượt thiếu reference qua biên dịch và sample, và một lượt do remote judge. Cách dừng trước metric là phù hợp: thiếu bằng chứng không đồng nghĩa chất lượng bằng 0 và không nên bị trộn với WA/RE/TLE của submission.")
    doc.add_heading("6.8 Đối chiếu thận trọng với CodeContests-O", level=2)
    add_table(doc, ["Hệ thống", "TPR", "TNR", "Quy mô/protocol"], [
        ["Final Testcase Generator", pct(s["macro_tpr"],2), pct(s["macro_tnr"],2), f"Macro trên {len(s['measured'])} bài; oracle hiện có"],
        ["CodeContests-O iter 3 [11]", "89,37%", "90,89%", "Khoảng 11 triệu lời giải; dataset và protocol riêng"],
        ["CodeContests+ 1× [11]", "85,68%", "86,03%", "Số liệu tái báo cáo trong CodeContests-O"],
        ["CodeContests [11]", "85,18%", "81,52%", "Số liệu tái báo cáo trong CodeContests-O"],
    ], widths=[4.5,2.4,2.4,6.5], font_size=8.8)
    p(doc,
      "TPR của prototype gần CodeContests-O, còn TNR thấp hơn khoảng 5 điểm phần trăm nếu chỉ nhìn hai con số. Không thể suy ra chênh lệch hiệu năng trực tiếp vì cách chọn bài, quy mô oracle, checker, ngôn ngữ và ngân sách khác nhau. Giá trị của đối chiếu nằm ở việc dùng cùng hai trục fidelity–discriminability và xác nhận rằng refinement qua feedback là hướng nghiên cứu hợp lý.")
    doc.add_heading("6.9 Trả lời câu hỏi nghiên cứu", level=2)
    add_table(doc, ["RQ", "Kết luận từ bằng chứng hiện tại"], [
        ["RQ1", f"Có: {s['ever_success_count']} suite khóa và {num(s['cases_total'])} test; nhưng 17/81 lượt hoàn tất batch vẫn thất bại."],
        ["RQ2", f"Macro TPR {pct(s['macro_tpr'],2)}, macro TNR {pct(s['macro_tnr'],2)}; 40/59 bài đồng thời đạt ít nhất 0,8."],
        ["RQ3", "CodeNet và Codeforces gần nhau ở TPR; CodeNet nhỉnh nhẹ ở TNR. Sai khác theo bài lớn hơn sai khác trung bình theo nguồn."],
        ["RQ4", f"36/57 run không vượt 180 giây; trung vị {num(s['generation_median'])} giây nhưng có phần đuôi dài."],
        ["RQ5", "Hai nút thắt lớn nhất là sinh đủ AC/validator và tạo đủ mutant WA được chứng nhận; nhãn lịch sử và checker là rủi ro đối với kết luận."],
    ], widths=[1.6,14.2], font_size=9.4)

    chapter(doc, "CHƯƠNG 7. GIỚI HẠN, HƯỚNG PHÁT TRIỂN VÀ KẾT LUẬN")
    doc.add_heading("7.1 Nguy cơ đối với độ giá trị", level=2)
    p(doc, "Độ giá trị cấu trúc. TPR/TNR dùng nhãn lịch sử, nên đo mức phù hợp với verdict nguồn hơn là chân lý toán học. Accepted có thể vẫn buggy [12]; Wrong Answer có thể phụ thuộc checker, runtime hoặc test cũ. Coverage là obligation coverage do LLM sinh, không phải structural code coverage.", "Độ giá trị cấu trúc.")
    p(doc, "Độ giá trị nội tại. Reference consensus giảm rủi ro nhưng không loại trừ việc nhiều reference cùng sai. Candidate và test cùng do LLM tạo có thể chia sẻ thiên lệch. Holdout chỉ có ở 5/60 run nên khả năng tổng quát nội bộ chưa được đo đồng đều.", "Độ giá trị nội tại.")
    p(doc, "Độ giá trị bên ngoài. 59 bài đo đủ chưa đại diện cho toàn bộ CodeNet hoặc Codeforces. Nhóm bài, checker và ngôn ngữ không được lấy mẫu ngẫu nhiên. Kết quả không nên suy rộng sang bài tương tác, output-only hoặc bài có checker phức tạp chưa được hỗ trợ.", "Độ giá trị bên ngoài.")
    p(doc, "Độ tin cậy. Báo cáo dùng artifact đã có và không chạy lại pipeline. 26 chương trình trong báo cáo validation checker bổ sung tạo 116 execution và không có accepted program thất bại, nhưng đó không phải rejudge toàn bộ 2.519 source. Workspace chưa gắn kết quả với commit.", "Độ tin cậy.")
    doc.add_heading("7.2 Hướng phát triển", level=2)
    add_numbered(doc, [
        "Gắn mọi batch với commit, model identifier, prompt version, image digest và seed; tạo manifest thí nghiệm bất biến.",
        "Tăng holdout pool và áp dụng cùng tỷ lệ trên mọi bài; giữ holdout kín khỏi refinement.",
        "Tạo quy trình triage false positive tự động: checker audit, multi-reference disagreement, brute-force cho miền nhỏ và replay nhiều toolchain.",
        "Phân cụm WA survivor theo hành vi để sinh fault obligation cụ thể thay vì chỉ thêm test ngẫu nhiên.",
        "Cho phép số vòng refinement thích nghi theo survivor và latency budget, thay vì cố định hai vòng.",
        "Mở rộng special checker, output có nhiều đáp án và tolerance số thực với test kiểm chứng riêng.",
        "Báo cáo confidence interval theo bài và thực hiện ablation cho consensus, witness, feedback và coverage model.",
        "Tách chi phí LLM, compile, execute, certification và oracle để tối ưu phần gây đuôi latency.",
        "Đánh giá lại trên tập bài chọn trước và frozen oracle để giảm selection bias.",
    ])
    doc.add_heading("7.3 Kết luận", level=2)
    p(doc,
      "Khóa luận đã xây dựng được một pipeline sinh testcase tự động có chuỗi chứng nhận rõ ràng, vòng lặp feedback, nhiều backend thực thi và cơ chế khóa trước đánh giá ngoài. Kết quả hiện tại đủ chứng minh tính khả thi: 60 bài có suite khóa, 1.144 testcase được chứng nhận, macro TPR 88,07% và macro TNR 85,78% trên 59 bài đo đủ. Phần lớn run đạt mốc 180 giây và hai nguồn dữ liệu chính cho kết quả tương đối nhất quán.")
    p(doc,
      "Điểm mạnh của hệ thống nằm ở việc không chấp nhận trực tiếp output của LLM: reference phải qua sample, mutant phải có witness và replay, testcase phải hợp lệ và có consensus, suite phải khóa trước oracle. Điểm còn thiếu nằm ở độ ổn định giữa các bài, quy mô holdout, audit false positive và reproducibility theo commit. Các giới hạn này xác định trực tiếp kế hoạch thực nghiệm cuối để chuyển prototype thành kết quả KLTN có độ tin cậy cao hơn.")

    chapter(doc, "TÀI LIỆU THAM KHẢO")
    refs = [
        "[1] I. Sommerville, Software Engineering, 9th ed. Boston, MA, USA: Addison-Wesley, 2011.",
        "[2] L. Copeland, A Practitioner’s Guide to Software Test Design. Norwood, MA, USA: Artech House, 2004.",
        "[3] C. Kaner, J. Falk, and H. Q. Nguyen, Testing Computer Software, 2nd ed. New York, NY, USA: Wiley, 1999.",
        "[4] P. McMinn, “Search-based software test data generation: A survey,” Software Testing, Verification and Reliability, vol. 14, no. 2, pp. 105–156, 2004, doi: 10.1002/stvr.294.",
        "[5] Y. Jia and M. Harman, “An analysis and survey of the development of mutation testing,” IEEE Transactions on Software Engineering, vol. 37, no. 5, pp. 649–678, 2011, doi: 10.1109/TSE.2010.62.",
        "[6] C. Lemieux, J. P. Inala, S. K. Lahiri, and S. Sen, “CodaMosa: Escaping coverage plateaus in test generation with pre-trained large language models,” in Proc. 45th IEEE/ACM Int. Conf. Software Engineering (ICSE), 2023, doi: 10.1109/ICSE48619.2023.00085.",
        "[7] J. A. Pizzorno and E. D. Berger, “CoverUp: Coverage-guided LLM-based test generation,” arXiv:2403.16218, 2024.",
        "[8] R. Puri et al., “CodeNet: A large-scale AI for code dataset for learning a diversity of coding tasks,” arXiv:2105.12655, 2021.",
        "[9] Y. Li et al., “Competition-level code generation with AlphaCode,” Science, vol. 378, no. 6624, pp. 1092–1097, 2022, doi: 10.1126/science.abq1158.",
        "[10] Z. Wang, S. Liu, Y. Sun, H. Li, and K. Shen, “CodeContests+: High-quality test case generation for competitive programming,” arXiv:2506.05817, 2025.",
        "[11] J. Cai, J. Zhu, R. Sun, K. Zhao, D. Xue, M. Feng, W. Zhou, and H. Li, “CodeContests-O: Powering LLMs via feedback-driven iterative test case generation,” arXiv:2601.13682, 2026.",
        "[12] S. Xie, S. Xie, F. Zhu, Y. Ji, and W. Zuo, “Coding agents as test-suite auditors: Finding what official suites miss while approaching what they catch,” arXiv:2608.01715, 2026.",
        "[13] Open R1, “codeforces-submissions,” Hugging Face Datasets, 2026. [Online]. Available: https://huggingface.co/datasets/open-r1/codeforces-submissions. Accessed: Oct. 13, 2026.",
        "[14] Codeforces, “G1. Yunli’s Subarray Queries (easy version), Codeforces Round 2009,” 2024. [Online]. Available: https://codeforces.com/contest/2009/problem/G1. Accessed: Oct. 13, 2026.",
    ]
    for ref in refs:
        para = doc.add_paragraph()
        para.paragraph_format.first_line_indent = Cm(-0.8)
        para.paragraph_format.left_indent = Cm(0.8)
        para.paragraph_format.line_spacing = 1.08
        para.paragraph_format.space_after = Pt(3)
        run = para.add_run(ref)
        run.font.size = Pt(11)

    # Detailed appendices use narrower margins and smaller text.
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width = Cm(29.7)
    sec.page_height = Cm(21)
    sec.top_margin = Cm(1.5)
    sec.bottom_margin = Cm(1.4)
    sec.left_margin = Cm(1.5)
    sec.right_margin = Cm(1.5)
    sec.header.is_linked_to_previous = False
    sec.footer.is_linked_to_previous = False
    set_repeat_header(sec, "Phụ lục A — Kết quả theo từng bài")
    add_page_number(sec.footer.paragraphs[0])
    doc.add_heading("PHỤ LỤC A. KẾT QUẢ THÀNH CÔNG GẦN NHẤT THEO TỪNG BÀI", level=1)
    p(doc,
      "Mỗi dòng là run thành công gần nhất sau khi loại batch chưa hoàn tất. TPR và TNR là oracle metric theo nhãn; CE và lỗi hạ tầng không nằm trong mẫu số logic. Dấu gạch ngang nghĩa là thiếu metric.")
    rows = []
    for r in sorted(s["runs"], key=lambda x:x["problem"].lower()):
        rows.append([
            r["problem"], dataset_of(r["problem"]), r["run"], r["cases"], r["iterations"], pct(r["coverage"]),
            pct(r["critical_coverage"]), pct(r["oracle_tpr"]), pct(r["oracle_tnr"]), classify(r["oracle_tpr"],r["oracle_tnr"]),
            num(r["generation_seconds"]), "Có" if r["holdout"] else "Không",
        ])
    for start in range(0, len(rows), 22):
        if start:
            doc.add_page_break()
        add_table(doc, ["Bài","Nguồn","Run","Test","Vòng","Coverage","Critical","TPR","TNR","Phân loại","Giây","Holdout"],
                  rows[start:start+22], widths=[2.1,2.2,3.2,1.2,1.2,1.8,1.8,1.5,1.5,2.8,1.7,1.7], font_size=7.2)

    sec2 = doc.add_section(WD_SECTION.NEW_PAGE)
    sec2.orientation = WD_ORIENT.PORTRAIT
    sec2.page_width = Cm(21)
    sec2.page_height = Cm(29.7)
    sec2.top_margin = Cm(1.8)
    sec2.bottom_margin = Cm(1.6)
    sec2.left_margin = Cm(2.0)
    sec2.right_margin = Cm(1.8)
    sec2.header.is_linked_to_previous = False
    sec2.footer.is_linked_to_previous = False
    set_repeat_header(sec2, "Phụ lục B — Batch, lỗi và nguồn artifact")
    add_page_number(sec2.footer.paragraphs[0])
    doc.add_heading("PHỤ LỤC B. TRẠNG THÁI BATCH", level=1)
    add_table(doc, ["Batch","Completed/Total","OK","Lỗi","finished_at","Xử lý"], [[
        b["id"], f"{b['completed']}/{b['total']}", b["succeeded"], b["failed"], "Có" if b["finished_at"] else "Không", "Dùng" if b["complete"] else "Loại"
    ] for b in e["batches"]], widths=[3.5,2.8,1.6,1.6,2.4,2.1], font_size=8.8)
    doc.add_heading("B.1 Lần thử gần nhất bị lỗi", level=2)
    latest_fail = sorted([a for a in s["latest_attempts"] if a["status"] == "failed"], key=lambda a:a["problem"])
    add_table(doc, ["Bài","Batch","Giây","Nguyên nhân","Từng thành công"], [[
        a["problem"], a["batch"], num(a["duration_seconds"]), a["failure_reason"], "Có" if a["problem"] in e["latest_success"] else "Chưa"
    ] for a in latest_fail], widths=[2.3,3.2,1.8,6.1,2.4], font_size=8.4)
    doc.add_heading("B.2 Bản đồ artifact", level=2)
    add_table(doc, ["Nhóm","Đường dẫn chính","Vai trò"], [
        ["Mô tả","README.md; Architect.md; Specs.md","Thiết kế và cách vận hành"],
        ["Pipeline","main.py; testgen/pipeline.py","Composition root và orchestration"],
        ["Generator","testgen/ollama.py; claude.py; chatgpt.py; fallback.py","Backend LLM và fallback"],
        ["Judge","testgen/docker_judge.py; remote_judge.py; native_judge.py","Compile/run và sandbox"],
        ["Checker","testgen/checker.py; testgen/checkers/*","So sánh output"],
        ["Cấu hình","config/pipeline.json; judge.json; sandbox.json; runtimes.json","Tham số thực nghiệm"],
        ["Dữ liệu","data/problems/<id>/","Statement, metadata, checker, oracle"],
        ["Run","output/.../report.json; test-suite.lock.json","Metric và suite bất biến"],
        ["Batch","output/batches/<timestamp>/batch-report.json; logs/","Trạng thái và lỗi"],
        ["Validation","output/codenet-additional20-validation/report.json","Kiểm tra 20 bài bổ sung"],
    ], widths=[3.0,7.4,5.4], font_size=8.3)

    doc.add_heading("PHỤ LỤC C. KIỂM TRA CHẤT LƯỢNG TÀI LIỆU", level=1)
    add_bullets(doc, [
        "Mọi hình lấy từ bài báo đều ghi tác giả, số hình gốc và số trích dẫn IEEE ngay dưới hình.",
        "Mọi biểu đồ kết quả dự án ghi rõ được tổng hợp từ report.json hoặc log; không trình bày như hình của nguồn ngoài.",
        "Bốn batch dừng giữa chừng không đóng góp kết quả; lỗi KeyboardInterrupt/KeyboardInput không được xem là lỗi logic của pipeline.",
        "Các so sánh với CodeContests-O được đánh dấu là tham khảo bối cảnh vì protocol không tương đương.",
        "Số liệu tổng hợp và bảng theo bài được sinh từ cùng một snapshot trong lần dựng tài liệu.",
    ])

    core = doc.core_properties
    core.title = "Hệ thống sinh và chứng nhận testcase tự động cho bài toán lập trình"
    core.subject = "Báo cáo khóa luận tốt nghiệp — Final Testcase Generator"
    core.author = "Nhóm thực hiện khóa luận"
    core.keywords = "testcase generation, LLM, mutation testing, CodeNet, Codeforces, TPR, TNR"
    core.comments = "Tổng hợp từ code, dữ liệu, log và tài liệu tham khảo đến ngày 13/10/2026."
    doc.save(OUTPUT)


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    evidence = old.load_evidence()
    # G1 is Codeforces contest 2009, problem G1, crawled manually. The older
    # evidence loader falls back to "Khác" because the local id lacks "cf".
    for row in evidence["data_rows"]:
        if row["problem"] == "G1":
            row["dataset"] = "Codeforces"
    if "G1" in evidence["metadata"]:
        evidence["metadata"]["G1"]["dataset"] = "Codeforces"
    summary = old.summarize(evidence)
    charts = create_charts(evidence, summary)
    snapshot = {
        "captured_at": "2026-10-13T00:00:00+07:00",
        "selection": "latest successful locked run per problem; incomplete batches excluded",
        "summary": {k:v for k,v in summary.items() if k not in {"runs","measured","generation","complete_batches","excluded_batches","latest_attempts"}},
        "runs": summary["runs"],
        "latest_attempts": summary["latest_attempts"],
        "batches": evidence["batches"],
    }
    (WORK_DIR / "evidence-snapshot.json").write_text(json.dumps(snapshot, ensure_ascii=False, indent=2), encoding="utf-8")
    build_document(evidence, summary, charts)
    print(json.dumps({"output":str(OUTPUT),"bytes":OUTPUT.stat().st_size,"runs":summary["ever_success_count"],"measured":len(summary["measured"])}, ensure_ascii=False))


if __name__ == "__main__":
    main()

from __future__ import annotations

import json
import math
import re
import statistics
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
OUT_DIR = ROOT / "reports" / "2026-10-12"
WORK_DIR = OUT_DIR / ".progress-report-work"
OUTPUT = OUT_DIR / "Bao-cao-tien-do-Final-TestcaseGenerator-2026-10-12-v2.docx"

NAVY = "17365D"
PALE_BLUE = "EAF2F8"
PALE_GRAY = "F5F7F9"
MID_GRAY = "667085"
LIGHT_BORDER = "D9D9D9"
BLACK = "000000"
WHITE = "FFFFFF"
GREEN = "1B5E20"
AMBER = "8A4B08"
RED = "8B1E1E"


def read_json(path: Path):
    return json.loads(path.read_text(encoding="utf-8"))


def rel(path: Path | str) -> str:
    p = Path(path)
    try:
        return str(p.resolve().relative_to(ROOT.resolve())).replace("\\", "/")
    except ValueError:
        return str(p).replace("\\", "/")


def is_pipeline_report(data: dict) -> bool:
    return "problem_id" in data and "cases" in data and "test_suite_lock" in data


def failure_reason(log_path: str | None) -> str:
    if not log_path:
        return "Không có log"
    path = Path(log_path)
    if not path.is_file():
        return "Thiếu log"
    text = path.read_text(encoding="utf-8", errors="replace")
    if "KeyboardInterrupt" in text or "KeyboardInput" in text:
        return "Bị ngắt bằng bàn phím"
    if "Generation incomplete" in text:
        return "Không sinh đủ AC hoặc validator"
    if "minimum witness-verified wrong-candidate pool" in text or "minimum verified wrong-candidate pool" in text:
        return "Không đủ mutant WA được chứng nhận"
    if "reference candidates passed compilation and official samples" in text:
        return "Không đủ reference qua biên dịch và sample"
    if "Remote judge unavailable" in text:
        return "Remote judge không sẵn sàng"
    matches = re.findall(r"(?:RuntimeError|ValueError|TimeoutError|ConnectionError):\s*(.+)", text)
    if matches:
        return matches[-1][:160]
    return "Lỗi khác có log"


def classify(tpr, tnr) -> str:
    if tpr is None or tnr is None:
        return "Chưa đo đủ"
    if tpr >= 0.9 and tnr >= 0.9:
        return "Mạnh"
    if tpr >= 0.8 and tnr >= 0.8:
        return "Đạt 0,8"
    return "Cần cải thiện"


def pct(value) -> str:
    if value is None:
        return "—"
    return f"{value * 100:.1f}%"


def num(value, digits=1) -> str:
    if value is None:
        return "—"
    if isinstance(value, int):
        return f"{value:,}".replace(",", ".")
    return f"{value:.{digits}f}".replace(".", ",")


def load_evidence() -> dict:
    batch_state = {}
    batches = []
    complete_attempts = []
    for batch_dir in sorted((ROOT / "output" / "batches").glob("*")):
        report_path = batch_dir / "batch-report.json"
        if not report_path.is_file():
            continue
        data = read_json(report_path)
        complete = bool(data.get("finished_at")) and data.get("completed") == data.get("total")
        batch_state[batch_dir.name] = complete
        row = {
            "id": batch_dir.name,
            "path": rel(report_path),
            "started_at": data.get("started_at"),
            "finished_at": data.get("finished_at"),
            "total": data.get("total", 0),
            "completed": data.get("completed", 0),
            "succeeded": data.get("succeeded", 0),
            "failed": data.get("failed", 0),
            "complete": complete,
            "results": data.get("results", []),
        }
        batches.append(row)
        if complete:
            for result in data.get("results", []):
                stamp = datetime.fromisoformat(data["finished_at"]).timestamp()
                complete_attempts.append({
                    "problem": result.get("problem"),
                    "batch": batch_dir.name,
                    "status": result.get("status"),
                    "duration_seconds": result.get("duration_seconds"),
                    "report": result.get("report"),
                    "log": result.get("log"),
                    "failure_reason": failure_reason(result.get("log")) if result.get("status") == "failed" else "",
                    "timestamp": stamp,
                })

    reports = []
    for path in (ROOT / "output").rglob("report.json"):
        data = read_json(path)
        if not is_pipeline_report(data):
            continue
        match = re.search(r"output[\\/]batches[\\/]([^\\/]+)[\\/]", str(path))
        batch_id = match.group(1) if match else None
        if batch_id and not batch_state.get(batch_id, False):
            continue
        if not (path.parent / "test-suite.lock.json").is_file():
            continue
        metrics = data.get("oracle_metrics") or {}
        dev_metrics = data.get("llm_metrics") or {}
        latency = data.get("generation_latency") or {}
        cert = data.get("mutant_certification") or {}
        reports.append({
            "problem": data.get("problem_id"),
            "path": rel(path),
            "run": batch_id or path.parent.name,
            "batch": batch_id,
            "timestamp": path.stat().st_mtime,
            "cases": len(data.get("cases") or []),
            "quarantined": len(data.get("quarantined") or []),
            "iterations": data.get("iterations_used"),
            "stop_reason": data.get("stop_reason"),
            "coverage": (data.get("coverage") or {}).get("coverage_ratio"),
            "critical_coverage": (data.get("coverage") or {}).get("critical_ratio"),
            "normal_coverage": (data.get("coverage") or {}).get("normal_ratio"),
            "oracle_tpr": metrics.get("tpr_label_based"),
            "oracle_tnr": metrics.get("tnr_label_based"),
            "oracle_positive_pass": metrics.get("positive_pass"),
            "oracle_positive_evaluable": metrics.get("positive_evaluable"),
            "oracle_negative_killed": metrics.get("negative_killed_logic"),
            "oracle_negative_evaluable": metrics.get("negative_evaluable"),
            "dev_tpr": dev_metrics.get("tpr_label_based"),
            "dev_tnr": dev_metrics.get("tnr_label_based"),
            "holdout": data.get("llm_holdout_metrics") is not None,
            "generation_seconds": latency.get("elapsed_seconds"),
            "generation_target_met": latency.get("target_met"),
            "mutant_stage": cert.get("stage_status"),
            "validated_mutants": cert.get("validated_wrong_candidates"),
        })

    latest_success = {}
    for report in reports:
        current = latest_success.get(report["problem"])
        if current is None or report["timestamp"] > current["timestamp"]:
            latest_success[report["problem"]] = report

    latest_attempt = {}
    for attempt in complete_attempts:
        current = latest_attempt.get(attempt["problem"])
        if current is None or attempt["timestamp"] > current["timestamp"]:
            latest_attempt[attempt["problem"]] = attempt
    for report in reports:
        if report["batch"] is None:
            attempt = {
                "problem": report["problem"], "batch": report["run"], "status": "succeeded",
                "duration_seconds": None, "report": report["path"], "log": None,
                "failure_reason": "", "timestamp": report["timestamp"],
            }
            current = latest_attempt.get(report["problem"])
            if current is None or attempt["timestamp"] > current["timestamp"]:
                latest_attempt[report["problem"]] = attempt

    metadata = {}
    data_rows = []
    for problem_dir in sorted((ROOT / "data" / "problems").iterdir()):
        if not problem_dir.is_dir():
            continue
        meta_path = problem_dir / "metadata.json"
        meta = read_json(meta_path) if meta_path.is_file() else {}
        checker_path = problem_dir / "checker.json"
        checker = read_json(checker_path) if checker_path.is_file() else {}
        if problem_dir.name.startswith("p") and problem_dir.name[1:].isdigit():
            dataset = "CodeNet"
        elif problem_dir.name.lower().startswith("cf"):
            dataset = "Codeforces"
        else:
            dataset = meta.get("dataset") or "Khác"
        row = {
            "problem": problem_dir.name,
            "title": meta.get("title", ""),
            "dataset": dataset,
            "statement": (problem_dir / "statement.md").is_file(),
            "metadata": meta_path.is_file(),
            "checker": checker.get("kind"),
            "oracle": (problem_dir / "oracle" / "manifest.json").is_file(),
        }
        data_rows.append(row)
        metadata[problem_dir.name] = row

    manifest_entries = []
    for path in (ROOT / "data" / "problems").rglob("manifest.json"):
        raw = read_json(path)
        if isinstance(raw, list):
            manifest_entries.extend(raw)

    validation_path = ROOT / "output" / "codenet-additional20-validation" / "report.json"
    validation = read_json(validation_path) if validation_path.is_file() else None

    return {
        "batches": batches,
        "complete_attempts": complete_attempts,
        "latest_attempt": latest_attempt,
        "reports": reports,
        "latest_success": latest_success,
        "metadata": metadata,
        "data_rows": data_rows,
        "manifest_entries": manifest_entries,
        "validation": validation,
        "validation_path": rel(validation_path),
    }


def summarize(e: dict) -> dict:
    runs = list(e["latest_success"].values())
    measured = [r for r in runs if r["oracle_tpr"] is not None and r["oracle_tnr"] is not None]
    generation = [r["generation_seconds"] for r in runs if r["generation_seconds"] is not None]
    pos_pass = sum(r["oracle_positive_pass"] or 0 for r in measured)
    pos_eval = sum(r["oracle_positive_evaluable"] or 0 for r in measured)
    neg_killed = sum(r["oracle_negative_killed"] or 0 for r in measured)
    neg_eval = sum(r["oracle_negative_evaluable"] or 0 for r in measured)
    complete_batches = [b for b in e["batches"] if b["complete"]]
    excluded_batches = [b for b in e["batches"] if not b["complete"]]
    latest_attempts = list(e["latest_attempt"].values())
    manifest = e["manifest_entries"]
    return {
        "runs": runs,
        "measured": measured,
        "generation": generation,
        "complete_batches": complete_batches,
        "excluded_batches": excluded_batches,
        "latest_attempts": latest_attempts,
        "complete_batch_attempts": len(e["complete_attempts"]),
        "complete_batch_success": sum(a["status"] == "succeeded" for a in e["complete_attempts"]),
        "complete_batch_failed": sum(a["status"] == "failed" for a in e["complete_attempts"]),
        "latest_success_count": sum(a["status"] == "succeeded" for a in latest_attempts),
        "latest_failed_count": sum(a["status"] == "failed" for a in latest_attempts),
        "ever_success_count": len(runs),
        "never_success": sorted(a["problem"] for a in latest_attempts if a["status"] == "failed" and a["problem"] not in e["latest_success"]),
        "cases_total": sum(r["cases"] for r in runs),
        "cases_avg": statistics.mean(r["cases"] for r in runs),
        "coverage_full": sum(r["coverage"] == 1 for r in runs),
        "quality_stop": sum(r["stop_reason"] == "all_quality_criteria_met" for r in runs),
        "iteration_stop": sum(r["stop_reason"] == "iteration_limit" for r in runs),
        "both_08": sum(r["oracle_tpr"] >= 0.8 and r["oracle_tnr"] >= 0.8 for r in measured),
        "both_09": sum(r["oracle_tpr"] >= 0.9 and r["oracle_tnr"] >= 0.9 for r in measured),
        "macro_tpr": statistics.mean(r["oracle_tpr"] for r in measured),
        "macro_tnr": statistics.mean(r["oracle_tnr"] for r in measured),
        "micro_tpr": pos_pass / pos_eval,
        "micro_tnr": neg_killed / neg_eval,
        "pos_pass": pos_pass,
        "pos_eval": pos_eval,
        "neg_killed": neg_killed,
        "neg_eval": neg_eval,
        "generation_avg": statistics.mean(generation),
        "generation_median": statistics.median(generation),
        "generation_le_180": sum(x <= 180 for x in generation),
        "holdout_count": sum(r["holdout"] for r in runs),
        "problem_dirs": len(e["data_rows"]),
        "statements": sum(r["statement"] for r in e["data_rows"]),
        "metadata_count": sum(r["metadata"] for r in e["data_rows"]),
        "checker_count": sum(r["checker"] is not None for r in e["data_rows"]),
        "oracle_bundle_count": sum(r["oracle"] for r in e["data_rows"]),
        "manifest_count": sum(1 for r in e["data_rows"] if r["oracle"]),
        "manifest_entries": len(manifest),
        "manifest_accepted": sum(x.get("label") == "Accepted" for x in manifest),
        "manifest_wrong": sum(x.get("label") == "Wrong Answer" for x in manifest),
        "manifest_unique": len({x.get("sha256") for x in manifest if x.get("sha256")}),
    }


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=110, bottom=100, end=110) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=LIGHT_BORDER, size="6") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_repeat_header(section, text: str) -> None:
    header = section.header
    p = header.paragraphs[0]
    p.clear()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(MID_GRAY)


def add_page_number(paragraph) -> None:
    paragraph.clear()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    run.font.name = "Arial"
    run.font.size = Pt(8)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for name, size, before, after in (
        ("Title", 23, 0, 16),
        ("Heading 1", 15, 16, 7),
        ("Heading 2", 12.5, 12, 5),
        ("Heading 3", 11, 9, 4),
    ):
        style = styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    styles["Title"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_ppr = styles["Title"]._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)
    styles["Subtitle"].font.name = "Arial"
    styles["Subtitle"].font.size = Pt(12)
    styles["Subtitle"].font.color.rgb = RGBColor.from_string(MID_GRAY)
    styles["Subtitle"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_table(doc, headers, rows, widths=None, font_size=8.5, alignments=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    if widths:
        # python-docx only changes tcW when assigning cell.width.  Older Word
        # versions may then rebuild tblGrid differently on later pages, which
        # can shift a wide table outside the printable area.  Pin both the
        # table width and every grid column so all appendix chunks render with
        # the same geometry.
        total_twips = sum(Cm(width).twips for width in widths)
        tbl_width = table._tbl.tblPr.first_child_found_in("w:tblW")
        tbl_width.set(qn("w:type"), "dxa")
        tbl_width.set(qn("w:w"), str(total_twips))
        for grid_col, width in zip(table._tbl.tblGrid.gridCol_lst, widths):
            grid_col.set(qn("w:w"), str(Cm(width).twips))
    set_table_borders(table)
    hdr = table.rows[0]
    repeat_table_header(hdr)
    prevent_row_split(hdr)
    for i, value in enumerate(headers):
        cell = hdr.cells[i]
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        r = p.add_run(str(value))
        r.bold = True
        r.font.name = "Arial"
        r.font.size = Pt(font_size)
        r.font.color.rgb = RGBColor.from_string(WHITE)
        if widths:
            cell.width = Cm(widths[i])
    for ridx, values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        for i, value in enumerate(values):
            cell = row.cells[i]
            if ridx % 2:
                set_cell_shading(cell, PALE_GRAY)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.alignment = alignments[i] if alignments else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(str(value))
            r.font.name = "Arial"
            r.font.size = Pt(font_size)
            if widths:
                cell.width = Cm(widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Cm(0.65)
        p.paragraph_format.first_line_indent = Cm(-0.25)
        p.add_run(item)


def add_numbered(doc, items):
    for index, item in enumerate(items, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.65)
        p.paragraph_format.first_line_indent = Cm(-0.5)
        p.paragraph_format.space_after = Pt(2)
        marker = p.add_run(f"{index}.  ")
        marker.font.name = "Arial"
        p.add_run(item)


def add_caption(doc, text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.italic = True
    r.font.name = "Arial"
    r.font.size = Pt(9)


def add_term(doc, term: str, explanation: str):
    p = doc.add_paragraph()
    p.paragraph_format.keep_together = True
    r = p.add_run(term + ". ")
    r.bold = True
    p.add_run(explanation)


def make_charts(e: dict, s: dict) -> dict:
    WORK_DIR.mkdir(parents=True, exist_ok=True)
    font_path = Path("C:/Windows/Fonts/arial.ttf")
    bold_path = Path("C:/Windows/Fonts/arialbd.ttf")
    font = ImageFont.truetype(str(font_path), 23)
    small = ImageFont.truetype(str(font_path), 19)
    bold = ImageFont.truetype(str(bold_path), 24)
    title_font = ImageFont.truetype(str(bold_path), 29)

    def canvas(title: str):
        image = Image.new("RGB", (1332, 666), "white")
        draw = ImageDraw.Draw(image)
        draw.text((666, 24), title, font=title_font, fill="#000000", anchor="ma")
        return image, draw

    def axes(draw):
        left, top, right, bottom = 115, 95, 1260, 570
        draw.line((left, top, left, bottom), fill="#666666", width=2)
        draw.line((left, bottom, right, bottom), fill="#666666", width=2)
        for i in range(6):
            y = bottom - (bottom - top) * i / 5
            draw.line((left, y, right, y), fill="#E1E5EA", width=1)
            draw.text((left - 16, y), f"{i/5:.1f}", font=small, fill="#555555", anchor="rm")
        return left, top, right, bottom

    groups = defaultdict(list)
    for run in s["runs"]:
        if run["problem"].startswith("p") and run["problem"][1:].isdigit():
            group = "CodeNet"
        elif run["problem"].lower().startswith("cf"):
            group = "Codeforces"
        else:
            group = "Khác"
        if run["oracle_tpr"] is not None and run["oracle_tnr"] is not None:
            groups[group].append(run)
    labels = [x for x in ("CodeNet", "Codeforces", "Khác") if groups.get(x)]
    tpr = [statistics.mean(r["oracle_tpr"] for r in groups[x]) for x in labels]
    tnr = [statistics.mean(r["oracle_tnr"] for r in groups[x]) for x in labels]
    image, draw = canvas("Macro TPR và TNR theo nhóm dữ liệu")
    left, top, right, bottom = axes(draw)
    group_width = (right - left) / len(labels)
    bar_width = min(120, group_width * 0.28)
    for i, label in enumerate(labels):
        center = left + group_width * (i + 0.5)
        for value, offset, color in ((tpr[i], -bar_width / 2, "#2F5597"), (tnr[i], bar_width / 2, "#70AD47")):
            x0 = center + offset - bar_width / 2
            x1 = center + offset + bar_width / 2
            y0 = bottom - value * (bottom - top)
            draw.rectangle((x0, y0, x1, bottom), fill=color)
            draw.text(((x0 + x1) / 2, y0 - 8), f"{value:.3f}", font=small, fill="#222222", anchor="ms")
        draw.text((center, bottom + 20), label, font=font, fill="#222222", anchor="ma")
    draw.rectangle((455, 615, 480, 635), fill="#2F5597")
    draw.text((490, 625), "TPR", font=small, fill="#222222", anchor="lm")
    draw.rectangle((690, 615, 715, 635), fill="#70AD47")
    draw.text((725, 625), "TNR", font=small, fill="#222222", anchor="lm")
    performance = WORK_DIR / "performance-by-dataset.png"
    image.save(performance)

    values = s["generation"]
    bins = [0, 60, 90, 120, 150, 180, 240, 360, 600, 900]
    counts = []
    for lo, hi in zip(bins[:-1], bins[1:]):
        counts.append(sum(lo <= value < hi for value in values))
    image, draw = canvas("Phân bố thời gian sinh và chứng nhận bộ test")
    left, top, right, bottom = 115, 95, 1260, 570
    draw.line((left, top, left, bottom), fill="#666666", width=2)
    draw.line((left, bottom, right, bottom), fill="#666666", width=2)
    max_count = max(counts) or 1
    step = max(1, math.ceil(max_count / 5))
    grid_max = step * 5
    for i in range(6):
        y = bottom - (bottom - top) * i / 5
        draw.line((left, y, right, y), fill="#E1E5EA", width=1)
        draw.text((left - 16, y), str(step * i), font=small, fill="#555555", anchor="rm")
    slot = (right - left) / len(counts)
    for i, count in enumerate(counts):
        x0 = left + i * slot + 8
        x1 = left + (i + 1) * slot - 8
        y0 = bottom - (count / grid_max) * (bottom - top)
        draw.rectangle((x0, y0, x1, bottom), fill="#5B9BD5")
        draw.text(((x0 + x1) / 2, y0 - 7), str(count), font=small, fill="#222222", anchor="ms")
        label = f"{bins[i]}–{bins[i+1]}"
        draw.text(((x0 + x1) / 2, bottom + 18), label, font=small, fill="#333333", anchor="ma")
    # The x-axis displays nine categorical time intervals with equal visual
    # width.  Therefore 180 seconds is the boundary after the fifth interval,
    # not 180/900 of the plot width (which would only be valid on a continuous
    # numeric axis).
    x180 = left + 5 * slot
    for y in range(top, bottom, 14):
        draw.line((x180, y, x180, min(y + 7, bottom)), fill="#C00000", width=2)
    draw.text((x180 + 8, top + 8), "Mốc 180 giây", font=small, fill="#C00000", anchor="la")
    draw.text(((left + right) / 2, 635), "Khoảng thời gian giây", font=font, fill="#222222", anchor="ma")
    latency = WORK_DIR / "generation-latency.png"
    image.save(latency)
    return {"performance": performance, "latency": latency}


def build_document(e: dict, s: dict, charts: dict) -> None:
    doc = Document()
    configure_styles(doc)
    doc.settings.odd_and_even_pages_header_footer = False
    section = doc.sections[0]
    section.top_margin = Cm(2.1)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.35)
    section.right_margin = Cm(2.1)
    section.different_first_page_header_footer = True
    set_repeat_header(section, "Báo cáo tiến độ Final Testcase Generator")
    add_page_number(section.footer.paragraphs[0])
    section.first_page_header.paragraphs[0].clear()
    section.first_page_footer.paragraphs[0].clear()

    p = doc.add_paragraph(style="Title")
    p.add_run("Báo cáo tiến độ hệ thống sinh testcase tự động")
    p = doc.add_paragraph(style="Subtitle")
    p.add_run("Dự án Final Testcase Generator")
    p = doc.add_paragraph(style="Subtitle")
    p.add_run("Báo cáo gửi giảng viên hướng dẫn")
    doc.add_paragraph()
    info = add_table(
        doc,
        ["Nội dung", "Thông tin"],
        [
            ["Thời điểm chốt dữ liệu", "17 giờ 07 phút ngày 12 tháng 10 năm 2026"],
            ["Phạm vi", "Code, dữ liệu và log hiện có trong Final_TestcaseGenerator"],
            ["Quy tắc chọn run", "Lấy kết quả hoàn tất gần nhất; loại toàn bộ batch bị dừng giữa chừng"],
            ["Mục đích", "Trình bày tiến độ, kết quả thực nghiệm, giới hạn và công việc tiếp theo"],
        ],
        widths=[4.2, 11.5],
        font_size=9.5,
        alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT],
    )
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Kết luận chính")
    r.bold = True
    r.font.size = Pt(12)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.add_run(
        f"Pipeline đã tạo và khóa bộ test thành công cho {s['ever_success_count']} bài. "
        f"Trong đó {len(s['measured'])} bài có đủ TPR và TNR từ tập oracle độc lập. "
        f"Macro TPR đạt {pct(s['macro_tpr'])}, macro TNR đạt {pct(s['macro_tnr'])}; "
        f"{s['both_08']}/{len(s['measured'])} bài đạt đồng thời hai chỉ số từ 0,8 và "
        f"{s['both_09']}/{len(s['measured'])} bài đạt đồng thời từ 0,9. "
        "Kết quả cho thấy pipeline đã vận hành được trên nhiều bài và nhiều ngôn ngữ, nhưng chất lượng chưa đồng đều giữa các bài."
    )
    doc.add_page_break()

    doc.add_heading("Nội dung báo cáo", level=1)
    add_numbered(doc, [
        "Phạm vi và phương pháp kiểm kê",
        "Thuật ngữ và cách đọc kết quả",
        "Kiến trúc codebase và luồng pipeline",
        "Dữ liệu thực nghiệm",
        "Giao thức đánh giá",
        "Kết quả đã đạt được",
        "Lỗi quan sát được và giới hạn",
        "Phần việc đã hoàn thành và hướng tiếp theo",
        "Phụ lục kết quả theo từng bài",
    ])
    p = doc.add_paragraph()
    p.add_run("Nguồn chính: ").bold = True
    p.add_run(
        "README.md, Architect.md, Specs.md, testgen/, config/, data/problems/, "
        "output/**/report.json, output/batches/*/batch-report.json và log của từng bài."
    )

    doc.add_heading("1 Phạm vi và phương pháp kiểm kê", level=1)
    doc.add_paragraph(
        "Báo cáo phản ánh trạng thái hiện có trên máy tại thời điểm chốt dữ liệu. Phạm vi chỉ gồm dự án "
        "Final_TestcaseGenerator. Không dùng dữ liệu bổ sung từ VPS, Kaggle hoặc kho khác và không chạy lại LLM, "
        "Docker, oracle hay bộ unit test trong lần lập báo cáo này. Các số liệu được đối chiếu từ mã nguồn, cấu hình, "
        "tệp dữ liệu, report JSON, khóa bộ test và log đã lưu."
    )
    doc.add_heading("1.1 Tiêu chí xác định kết quả hoàn tất", level=2)
    add_bullets(doc, [
        "Batch chỉ được đưa vào thống kê khi có finished_at và completed bằng total trong batch-report.json.",
        "Run thành công phải có report.json đúng schema pipeline và có test-suite.lock.json trong cùng thư mục.",
        "Kết quả theo bài lấy run thành công gần nhất trong tập hợp hợp lệ. Nếu lần thử mới nhất thất bại nhưng bài từng thành công trước đó, báo cáo giữ cả trạng thái thất bại mới nhất và số liệu thành công gần nhất.",
        "Bốn batch chưa hoàn tất bị loại toàn bộ, kể cả khi bên trong có một vài bài đã sinh report. Không có log nào chứa KeyboardInterrupt hoặc KeyboardInput trong 115 log batch đã kiểm tra.",
        "Report kiểm tra checker bổ sung được trình bày riêng vì không dùng schema của một run pipeline thông thường.",
    ])
    doc.add_heading("1.2 Quy mô bằng chứng", level=2)
    add_table(
        doc,
        ["Hạng mục", "Giá trị", "Ý nghĩa"],
        [
            ["Batch hoàn tất", len(s["complete_batches"]), "Được dùng để thống kê"],
            ["Batch bị loại", len(s["excluded_batches"]), "Dừng giữa chừng hoặc thiếu finished_at"],
            ["Lượt bài trong batch hoàn tất", s["complete_batch_attempts"], f"{s['complete_batch_success']} thành công và {s['complete_batch_failed']} thất bại"],
            ["Bài đã được thử", len(s["latest_attempts"]), f"Lần thử gần nhất: {s['latest_success_count']} thành công, {s['latest_failed_count']} thất bại"],
            ["Bài từng có report hoàn chỉnh", s["ever_success_count"], "Có khóa bộ test"],
            ["Bài có đủ oracle TPR và TNR", len(s["measured"]), "Được dùng trong tổng hợp chất lượng"],
        ],
        widths=[5.0, 2.5, 8.2],
        font_size=9,
        alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT],
    )

    doc.add_heading("2 Thuật ngữ và cách đọc kết quả", level=1)
    doc.add_paragraph(
        "Phần này giải thích các thuật ngữ được dùng xuyên suốt báo cáo. Các khái niệm được diễn giải theo cách pipeline "
        "đang triển khai, vì vậy người đọc có thể dùng trực tiếp để kiểm tra report JSON và log."
    )
    doc.add_heading("2.1 Đầu vào và đối tượng đánh giá", level=2)
    add_term(doc, "Problem", "Một bài toán lập trình gồm mô tả, định dạng input, định dạng output, giới hạn và ví dụ chính thức. Mỗi problem có một mã như p02548 hoặc cf1004a.")
    add_term(doc, "Statement", "Tệp statement.md chứa đề bài chuẩn hóa. Đây là nguồn thông tin chính để LLM phân tích miền dữ liệu, sinh code và đề xuất testcase.")
    add_term(doc, "Candidate solution", "Chương trình ứng viên dùng để đánh giá bộ test. Pipeline sinh hai nhóm: candidate_ac được kỳ vọng đúng và candidate_wa được kỳ vọng sai.")
    add_term(doc, "Accepted và Wrong Answer", "Accepted là nhãn cho chương trình được nguồn dữ liệu ghi nhận là đúng. Wrong Answer là nhãn cho chương trình từng cho kết quả sai. Đây là nhãn ban đầu, không phải chân lý tuyệt đối trong môi trường hiện tại.")
    add_term(doc, "Mutant", "Một lời giải sai được tạo bằng cách đưa lỗi có chủ đích vào lời giải đúng hoặc sinh trực tiếp theo một lớp lỗi. Mutant chỉ được dùng để tính TNR sau khi vượt qua bước chứng nhận.")
    add_term(doc, "Witness", "Một input cụ thể được kỳ vọng làm mutant sai. Witness phải hợp lệ, được validator chấp nhận, có expected output đáng tin cậy và tái hiện lỗi ổn định.")

    doc.add_heading("2.2 Thành phần xác định tính hợp lệ và đúng sai", level=2)
    add_term(doc, "Validator", "Chương trình kiểm tra input có đúng grammar và ràng buộc của đề hay không. Validator chỉ trả lời input hợp lệ; nó không xác định expected output.")
    add_term(doc, "Checker", "Thành phần so sánh output thực tế với expected output. Dự án có EXACT, TOKEN, CASE_INSENSITIVE_TOKEN, DECIMAL_EXACT, FLOAT và PYTHON_SCRIPT. Checker phải phù hợp với semantics của đề.")
    add_term(doc, "Reference solution", "Candidate AC được dùng để tính output tham chiếu. Pipeline cần ít nhất hai reference chạy thành công và đồng ý theo checker trước khi chứng nhận một testcase.")
    add_term(doc, "Reference consensus", "Sự đồng thuận của các reference về expected output. Đồng thuận làm giảm rủi ro dùng một lời giải duy nhất, nhưng nhiều reference vẫn có thể cùng sai.")
    add_term(doc, "Certified testcase", "Testcase đã qua validator, không trùng, có expected output từ sample hoặc reference consensus, và đủ bằng chứng để đưa vào test suite.")
    add_term(doc, "Quarantine", "Nơi lưu đề xuất bị từ chối do input sai, reference bất đồng, sample không khớp hoặc nguyên nhân khác. Test trong quarantine không được tính là testcase hợp lệ.")

    doc.add_heading("2.3 Các tập dùng trong thí nghiệm", level=2)
    add_term(doc, "Development candidates", "Nhóm candidate được chấm trong các vòng sinh test. Kết quả của nhóm này có thể được đưa vào feedback cho vòng sau.")
    add_term(doc, "Holdout candidates", "Nhóm mutant được giữ kín khỏi prompt sinh test. Nhóm này chỉ được chấm sau khi bộ test đã khóa, giúp kiểm tra khả năng tổng quát trong nội bộ pool LLM.")
    add_term(doc, "Oracle", "Tập chương trình bên ngoài gồm Accepted và Wrong Answer từ CodeNet hoặc Codeforces. Oracle được chấm sau khi khóa test suite và không được đưa ngược vào prompt.")
    add_term(doc, "Test suite lock", "Tệp test-suite.lock.json lưu hash và danh sách bộ test cuối. Pipeline kiểm tra khóa trước và sau khi chấm oracle để tránh thay đổi test theo kết quả ngoài.")

    doc.add_heading("2.4 Verdict", level=2)
    add_table(
        doc,
        ["Verdict", "Diễn giải", "Cách dùng trong metric"],
        [
            ["PASS", "Chương trình vượt toàn bộ testcase", "Tính là giữ được một nhãn AC"],
            ["WA", "Output không được checker chấp nhận", "Tính là bắt được một nhãn WA"],
            ["RE", "Chương trình lỗi khi chạy", "Tính là bắt được một nhãn WA"],
            ["TLE", "Vượt giới hạn thời gian", "Tính là bắt được nhãn WA, nhưng phản ánh cả hiệu năng"],
            ["CE", "Không biên dịch được", "Không đưa vào mẫu số logic"],
            ["INFRA", "Lỗi Docker, SSH hoặc hạ tầng", "Không đưa vào mẫu số; pipeline retry riêng"],
        ],
        widths=[2.2, 7.0, 6.5],
        font_size=9,
        alignments=[WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT],
    )

    doc.add_heading("2.5 TPR TNR và coverage", level=2)
    add_term(doc, "TPR label based", "Tỷ lệ chương trình mang nhãn Accepted hoặc candidate_ac vẫn PASS trên bộ test, sau khi loại CE và INFRA khỏi mẫu số. Chỉ số cao cho thấy bộ test ít loại nhầm chương trình đúng theo nhãn.")
    add_term(doc, "TNR label based", "Tỷ lệ chương trình mang nhãn Wrong Answer hoặc candidate_wa bị WA, RE hoặc TLE, sau khi loại CE và INFRA. Chỉ số cao cho thấy bộ test bắt được nhiều chương trình sai theo nhãn.")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("TPR = số AC PASS / số AC đánh giá được     TNR = số WA bị bắt / số WA đánh giá được")
    r.bold = True
    r.font.size = Pt(10)
    add_term(doc, "Macro average", "Tính TPR hoặc TNR riêng cho từng bài rồi lấy trung bình. Mỗi bài có trọng số bằng nhau.")
    add_term(doc, "Micro average", "Cộng số chương trình PASS, bị bắt và mẫu số của tất cả bài trước khi chia. Bài có nhiều chương trình đánh giá được sẽ có trọng số lớn hơn.")
    add_term(doc, "Coverage obligation", "Một yêu cầu phủ được sinh từ test model, chẳng hạn biên nhỏ nhất, biên lớn nhất, quan hệ đặc biệt, trường hợp vô nghiệm hoặc stress. Coverage ratio là tỷ lệ obligation được gắn với ít nhất một testcase được chứng nhận.")
    add_term(doc, "Critical coverage", "Tỷ lệ obligation ưu tiên critical đã được phủ. Cấu hình hiện tại yêu cầu 100 phần trăm để dừng sớm theo chất lượng.")
    add_term(doc, "Normal coverage", "Tỷ lệ obligation ưu tiên normal đã được phủ. Cấu hình hiện tại đặt mục tiêu 90 phần trăm.")
    add_term(doc, "Label based và provisional", "TPR và TNR đang đo mức phù hợp với nhãn lịch sử. Một AC bị fail có thể do testcase sai, checker sai, khác toolchain hoặc nhãn sai; một WA sống sót có thể do bộ test còn yếu. Vì vậy metric chưa phải chứng minh hình thức về độ đúng.")

    doc.add_heading("2.6 Vòng lặp và điều kiện dừng", level=2)
    add_term(doc, "Baseline iteration", "Vòng đầu chỉ dùng statement và test model để sinh testcase phủ miền, chưa dùng source mutant hay kết quả oracle.")
    add_term(doc, "Adaptive iteration", "Vòng sau nhắm vào mutant còn sống và khoảng trống coverage. Feedback chỉ chứa dữ liệu phát triển đã giới hạn, không chứa holdout hoặc oracle.")
    add_term(doc, "Top up", "Lượt sinh bổ sung khi số testcase được chấp nhận thấp hơn quota do trùng, validator từ chối hoặc reference bất đồng.")
    add_term(doc, "All quality criteria met", "Run dừng vì đã đủ số vòng và số test tối thiểu, đạt ngưỡng metric phát triển, đạt coverage và hết tiến triển cần thiết.")
    add_term(doc, "Iteration limit", "Run đã dùng hết số vòng cấu hình mà chưa thỏa toàn bộ điều kiện dừng sớm. Run vẫn có thể có oracle metric tốt hoặc coverage cao; tên này không đồng nghĩa run hỏng.")

    doc.add_heading("3 Kiến trúc codebase và luồng pipeline", level=1)
    doc.add_paragraph(
        "Codebase tách domain, bộ sinh, judge, checker, nguồn oracle và orchestration. main.py là composition root: đọc CLI và cấu hình, "
        "chọn backend rồi tạo Pipeline. testgen/pipeline.py điều phối vòng đời của một bài. Cấu trúc này cho phép thay LLM, judge, checker "
        "hoặc nguồn oracle mà không đổi toàn bộ luồng."
    )
    add_table(
        doc,
        ["Thành phần", "Vai trò hiện tại"],
        [
            ["main.py", "CLI, batch runner, chọn LLM và judge, preflight hạ tầng"],
            ["testgen/domain.py", "Mô hình Solution, Case, Result và interface Judge, Generator, OracleSource"],
            ["testgen/ollama.py", "Sinh test model, AC, validator, mutant, testcase, repair và audit"],
            ["testgen/claude.py và chatgpt.py", "Backend LLM trực tiếp thay cho Ollama khi được chọn"],
            ["testgen/fallback.py", "Chuỗi fallback backend cho từng nhiệm vụ"],
            ["testgen/pipeline.py", "Chứng nhận reference và mutant, vòng sinh test, khóa suite và tính metric"],
            ["testgen/checker.py và checkers", "Registry và plugin so sánh output"],
            ["testgen/docker_judge.py", "Biên dịch và chạy submission trong Docker local"],
            ["testgen/remote_judge.py", "Chuyển công việc chấm sang KHOI qua SSH"],
            ["testgen/native_judge.py", "Backend smoke test Linux có cờ xác nhận vì không cách ly filesystem và network"],
            ["testgen/retrying_judge.py", "Retry riêng cho INFRA, không retry verdict của code"],
            ["testgen/codenet.py", "Nạp oracle cache hoặc xuất nguồn từ kho CodeNet"],
            ["testgen/coverage.py", "Chuẩn hóa obligation và tính coverage"],
            ["testgen/progress.py", "Ghi tiến độ và latency.jsonl"],
        ],
        widths=[5.0, 10.7],
        font_size=8.8,
        alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT],
    )
    doc.add_heading("3.1 Luồng xử lý một bài", level=2)
    add_numbered(doc, [
        "Đọc statement, metadata, checker profile và cấu hình runtime.",
        "Chạy preflight để chọn Docker local, remote judge hoặc native smoke test theo tham số.",
        "Sinh test model gồm grammar, constraints, quan hệ, fault model và coverage obligation.",
        "Sinh candidate AC, candidate WA và validator; lưu response và source để audit.",
        "Biên dịch candidate, kiểm tra AC trên sample và chọn reference đủ điều kiện.",
        "Chứng nhận mutant bằng sample, witness, validator, reference consensus và hai lần tái hiện lỗi.",
        "Sinh testcase baseline, sửa input bị validator từ chối, loại trùng và chứng nhận expected output.",
        "Chấm development candidates, cập nhật survivor và coverage rồi sinh vòng adaptive nếu cần.",
        "Khóa test suite và chấm holdout nếu pool đủ lớn.",
        "Tải hoặc đọc oracle cache, chấm bằng cùng test suite và ghi report.json.",
        "Khi chạy batch, ghi log từng bài và batch-report.json; lỗi một bài không làm mất kết quả bài đã hoàn tất."
    ])
    doc.add_heading("3.2 Cấu hình đang sử dụng", level=2)
    add_table(
        doc,
        ["Nhóm", "Giá trị chính"],
        [
            ["Ứng viên", "3 candidate AC, 5 candidate WA, tối thiểu 3 mutant được chứng nhận"],
            ["Sinh test", "10 test được chấp nhận mỗi vòng, tối thiểu 2 vòng và 18 test, tối đa 2 vòng theo config hiện tại"],
            ["Coverage", "100% critical và 90% normal cho điều kiện dừng sớm"],
            ["LLM mặc định", "gemma4:31b-cloud, fallback nemotron-3-ultra:cloud rồi Claude theo chế độ fallback"],
            ["Judge", "auto, 8 worker chạy, 2 worker compile, batch remote 8 input"],
            ["Sandbox", "512 MiB RAM, 1 CPU, 64 process, run timeout 5 giây, compile timeout 90 giây"],
            ["Cách ly", "Không mạng, root filesystem chỉ đọc, cap drop, no new privileges, tmpfs 32 MiB"],
        ],
        widths=[4.2, 11.5],
        font_size=9,
        alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT],
    )

    doc.add_heading("4 Dữ liệu thực nghiệm", level=1)
    doc.add_paragraph(
        f"Thư mục data/problems hiện có {s['problem_dirs']} bài: 40 mã CodeNet, 31 mã Codeforces và 1 bài thuộc nhóm khác. "
        f"Tất cả {s['statements']} thư mục có statement.md; {s['metadata_count']} có metadata.json; "
        f"{s['checker_count']} có checker profile riêng và {s['oracle_bundle_count']} có oracle manifest cục bộ."
    )
    checker_counts = Counter(r["checker"] for r in e["data_rows"] if r["checker"])
    add_table(
        doc,
        ["Loại checker profile", "Số bài", "Ý nghĩa"],
        [
            ["TOKEN", checker_counts.get("TOKEN", 0), "So sánh token theo thứ tự, bỏ khác biệt whitespace"],
            ["EXACT", checker_counts.get("EXACT", 0), "So sánh nội dung sau chuẩn hóa newline cuối"],
            ["PYTHON_SCRIPT", checker_counts.get("PYTHON_SCRIPT", 0), "Dùng checker theo semantics riêng của bài"],
            ["Không có profile riêng", s["problem_dirs"] - s["checker_count"], "Dùng default TOKEN và cần rà trước khi coi metric là verified"],
        ],
        widths=[4.2, 2.2, 9.3],
        font_size=9,
        alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT],
    )
    languages = Counter(x.get("language") for x in e["manifest_entries"])
    add_table(
        doc,
        ["Hạng mục oracle", "Số lượng"],
        [
            ["Oracle manifest", s["manifest_count"]],
            ["Source có hash riêng", s["manifest_unique"]],
            ["Nhãn Accepted", s["manifest_accepted"]],
            ["Nhãn Wrong Answer", s["manifest_wrong"]],
            ["Python", languages.get("Python", 0)],
            ["C++", languages.get("C++", 0)],
            ["C", languages.get("C", 0)],
            ["Java", languages.get("Java", 0)],
        ],
        widths=[8.5, 4.0],
        font_size=9,
        alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER],
    )
    doc.add_paragraph(
        "Mỗi manifest ghi source tương đối, dialect, nhãn, SHA 256 và thông tin nguồn. Tổng số Accepted và Wrong Answer gần cân bằng, "
        "giúp TPR và TNR không bị chênh mạnh chỉ vì quy mô hai lớp. Tuy nhiên từng run thường dùng giới hạn 20 chương trình mỗi nhãn."
    )
    if e["validation"]:
        validation = e["validation"]
        doc.add_heading("4.1 Kiểm tra độc lập checker của nhóm CodeNet bổ sung", level=2)
        doc.add_paragraph(
            f"Artifact {e['validation_path']} ghi nhận {validation.get('programs')} chương trình được chạy trên "
            f"{validation.get('executions')} cặp chương trình và testcase, bao phủ "
            f"{len({x.get('problem_id') for x in validation.get('results', [])})} bài. "
            f"Số chương trình Accepted thất bại là {validation.get('accepted_program_failures')}. "
            "Phạm vi này dùng sample chính thức và probe viết tay cho special checker; đây không phải rejudge toàn bộ 800 source."
        )

    doc.add_heading("5 Giao thức đánh giá", level=1)
    doc.add_paragraph(
        "Hai lớp đánh giá được giữ tách biệt. Lớp phát triển dùng candidate do LLM sinh để dẫn dắt vòng adaptive. Lớp bên ngoài dùng oracle "
        "chỉ sau khi test suite đã khóa. Report lưu verdict từng source để có thể kiểm tra trường hợp metric thấp thay vì chỉ nhìn một tỷ lệ tổng."
    )
    add_table(
        doc,
        ["Lớp", "Thời điểm chấm", "Có vào feedback", "Mục đích"],
        [
            ["Development", "Sau từng vòng", "Có", "Diệt mutant đã biết và tìm coverage gap"],
            ["Holdout", "Sau suite lock", "Không", "Kiểm tra tổng quát trên mutant LLM được giữ kín"],
            ["Oracle", "Sau suite lock", "Không", "Đánh giá ngoài bằng source lịch sử độc lập"],
        ],
        widths=[3.0, 3.4, 3.0, 6.3],
        font_size=9,
        alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT],
    )
    doc.add_paragraph(
        "Các kết quả trình bày ở phần sau ưu tiên oracle metric. CE và INFRA bị loại khỏi mẫu số; TLE được tính là bắt được WA nhưng cần đọc cùng verdict "
        "nếu muốn tách lỗi hiệu năng khỏi lỗi logic. Macro và micro đều được báo để tránh kết luận dựa trên một cách gộp duy nhất."
    )

    doc.add_heading("6 Kết quả đã đạt được", level=1)
    doc.add_heading("6.1 Tổng quan định lượng", level=2)
    add_table(
        doc,
        ["Chỉ số", "Kết quả", "Cách hiểu"],
        [
            ["Bài có report hoàn chỉnh", s["ever_success_count"], "Có test suite lock"],
            ["Testcase được chứng nhận", int(s["cases_total"]), f"Trung bình {num(s['cases_avg'], 2)} test mỗi bài"],
            ["Coverage đạt 100%", f"{s['coverage_full']}/{s['ever_success_count']}", "Theo obligation trong test model"],
            ["Dừng vì đủ chất lượng", f"{s['quality_stop']}/{s['ever_success_count']}", "Thỏa toàn bộ điều kiện dừng sớm"],
            ["Dừng do hết vòng", f"{s['iteration_stop']}/{s['ever_success_count']}", "Có report nhưng còn ít nhất một điều kiện chưa đạt"],
            ["Macro TPR oracle", pct(s["macro_tpr"]), f"Trên {len(s['measured'])} bài đo đủ"],
            ["Macro TNR oracle", pct(s["macro_tnr"]), f"Trên {len(s['measured'])} bài đo đủ"],
            ["Micro TPR oracle", pct(s["micro_tpr"]), f"{s['pos_pass']}/{s['pos_eval']} AC PASS"],
            ["Micro TNR oracle", pct(s["micro_tnr"]), f"{s['neg_killed']}/{s['neg_eval']} WA bị bắt"],
            ["Cả TPR và TNR từ 0,8", f"{s['both_08']}/{len(s['measured'])}", "Ngưỡng theo báo cáo batch hiện có"],
            ["Cả TPR và TNR từ 0,9", f"{s['both_09']}/{len(s['measured'])}", "Nhóm kết quả mạnh"],
            ["Có đánh giá holdout", f"{s['holdout_count']}/{s['ever_success_count']}", "Chỉ xuất hiện khi pool mutant đủ lớn"],
        ],
        widths=[5.3, 3.0, 7.4],
        font_size=8.8,
        alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT],
    )

    doc.add_heading("6.2 Kết quả theo nguồn dữ liệu", level=2)
    group_rows = []
    for group_name, predicate in (
        ("CodeNet", lambda p: p.startswith("p") and p[1:].isdigit()),
        ("Codeforces", lambda p: p.lower().startswith("cf")),
        ("Khác", lambda p: not (p.startswith("p") and p[1:].isdigit()) and not p.lower().startswith("cf")),
    ):
        group = [r for r in s["runs"] if predicate(r["problem"])]
        measured = [r for r in group if r["oracle_tpr"] is not None and r["oracle_tnr"] is not None]
        group_rows.append([
            group_name,
            len(group),
            len(measured),
            pct(statistics.mean(r["oracle_tpr"] for r in measured)) if measured else "—",
            pct(statistics.mean(r["oracle_tnr"] for r in measured)) if measured else "—",
            f"{sum(r['oracle_tpr'] >= 0.8 and r['oracle_tnr'] >= 0.8 for r in measured)}/{len(measured)}" if measured else "—",
            f"{sum(r['coverage'] == 1 for r in group)}/{len(group)}" if group else "—",
        ])
    add_table(
        doc,
        ["Nhóm", "Run", "Đo đủ", "TPR TB", "TNR TB", "Cả hai từ 0,8", "Coverage 100%"],
        group_rows,
        widths=[3.0, 1.7, 1.8, 2.2, 2.2, 2.8, 2.8],
        font_size=8.5,
        alignments=[WD_ALIGN_PARAGRAPH.LEFT] + [WD_ALIGN_PARAGRAPH.CENTER] * 6,
    )
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(charts["performance"]), width=Inches(5.9))
    add_caption(doc, "Hình 1 Macro TPR và TNR theo nhóm dữ liệu")
    doc.add_paragraph(
        "CodeNet và Codeforces có macro TPR gần như bằng nhau. Macro TNR CodeNet nhỉnh hơn nhẹ. Nhóm Khác chỉ gồm một bài G1, "
        "nên không đủ để suy rộng; TNR thấp của bài này làm cột nhóm Khác giảm rõ rệt."
    )

    doc.add_heading("6.3 Thời gian sinh và chứng nhận bộ test", level=2)
    doc.add_paragraph(
        f"Có {len(s['generation'])} run thành công gần nhất ghi generation_latency. Trung vị là {num(s['generation_median'])} giây, "
        f"trung bình là {num(s['generation_avg'])} giây và {s['generation_le_180']}/{len(s['generation'])} run không vượt mốc 180 giây. "
        "Trung bình cao hơn trung vị vì một số run mất nhiều phút. generation_latency chỉ tính từ lúc bắt đầu đến khi có suite được chứng nhận và holdout; "
        "không bao gồm phần oracle bên ngoài."
    )
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(charts["latency"]), width=Inches(6.5))
    add_caption(doc, "Hình 2 Phân bố thời gian sinh và chứng nhận bộ test")
    slowest = sorted(
        (r for r in s["runs"] if r["generation_seconds"] is not None),
        key=lambda r: r["generation_seconds"], reverse=True,
    )[:8]
    add_table(
        doc,
        ["Bài", "Giây", "Test", "Coverage", "Dừng"],
        [[r["problem"], num(r["generation_seconds"]), r["cases"], pct(r["coverage"]), "Đủ chất lượng" if r["stop_reason"] == "all_quality_criteria_met" else "Hết vòng"] for r in slowest],
        widths=[3.2, 2.2, 2.0, 2.7, 4.2],
        font_size=8.8,
        alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER],
    )

    doc.add_heading("6.4 Những bài có chỉ số cần ưu tiên cải thiện", level=2)
    weakest = sorted(
        s["measured"],
        key=lambda r: (min(r["oracle_tpr"], r["oracle_tnr"]), (r["oracle_tpr"] + r["oracle_tnr"]) / 2),
    )[:12]
    add_table(
        doc,
        ["Bài", "TPR", "TNR", "Coverage", "Test", "Nhận xét"],
        [[
            r["problem"], pct(r["oracle_tpr"]), pct(r["oracle_tnr"]), pct(r["coverage"]), r["cases"],
            "Ưu tiên bảo toàn AC" if r["oracle_tpr"] < r["oracle_tnr"] else "Ưu tiên bắt WA",
        ] for r in weakest],
        widths=[2.7, 2.0, 2.0, 2.5, 1.6, 5.2],
        font_size=8.6,
        alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT],
    )
    doc.add_paragraph(
        "TPR thấp cần được điều tra theo thứ tự checker, expected output, toolchain và nhãn Accepted trước khi kết luận testcase quá chặt. "
        "TNR thấp cho thấy bộ test chưa phân biệt được đủ chương trình Wrong Answer hoặc các nhãn WA không còn tái hiện trong môi trường hiện tại."
    )

    doc.add_heading("6.5 Trạng thái batch", level=2)
    batch_rows = []
    for b in e["batches"]:
        state = "Đưa vào" if b["complete"] else "Loại"
        batch_rows.append([
            b["id"], f"{b['completed']}/{b['total']}", b["succeeded"], b["failed"],
            "Có" if b["finished_at"] else "Không", state,
        ])
    add_table(
        doc,
        ["Batch", "Hoàn tất", "OK", "Lỗi", "finished_at", "Xử lý"],
        batch_rows,
        widths=[3.6, 2.4, 1.7, 1.7, 2.2, 2.2],
        font_size=8.6,
        alignments=[WD_ALIGN_PARAGRAPH.LEFT] + [WD_ALIGN_PARAGRAPH.CENTER] * 5,
    )
    doc.add_paragraph(
        "Bốn batch bị loại là 20261010-160446, 20261010-162624, 20261012-160932 và 20261012-161357. "
        "Việc loại theo batch bảo đảm báo cáo không trộn kết quả từ một quá trình bị dừng giữa chừng."
    )

    doc.add_heading("7 Lỗi quan sát được và giới hạn", level=1)
    doc.add_heading("7.1 Lỗi trong các batch hoàn tất", level=2)
    complete_failures = [a for a in e["complete_attempts"] if a["status"] == "failed"]
    failure_counts = Counter(a["failure_reason"] for a in complete_failures)
    add_table(
        doc,
        ["Nhóm lỗi", "Số lượt", "Ý nghĩa"],
        [[
            reason, count,
            {
                "Không sinh đủ AC hoặc validator": "Dừng trước khi có đủ đầu vào tin cậy cho pipeline",
                "Không đủ mutant WA được chứng nhận": "Đề xuất WA không tái hiện ổn định hoặc không qua witness",
                "Không đủ reference qua biên dịch và sample": "Không thể thiết lập reference consensus an toàn",
                "Remote judge không sẵn sàng": "Lỗi SSH hoặc hạ tầng, không được gán thành verdict submission",
            }.get(reason, "Xem log chi tiết"),
        ] for reason, count in failure_counts.most_common()],
        widths=[6.5, 2.0, 7.2],
        font_size=8.8,
        alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT],
    )
    doc.add_paragraph(
        f"Trong {s['complete_batch_attempts']} lượt thuộc batch hoàn tất, có {s['complete_batch_failed']} lượt dừng trước metric. "
        "Pipeline không gán các lượt này TPR hoặc TNR bằng 0. Đây là lựa chọn đúng về mặt đo lường vì thiếu report không đồng nghĩa bộ test có chất lượng bằng 0."
    )
    doc.add_heading("7.2 Lần thử gần nhất bị lỗi", level=2)
    latest_failures = sorted((a for a in s["latest_attempts"] if a["status"] == "failed"), key=lambda a: a["problem"])
    add_table(
        doc,
        ["Bài", "Batch", "Thời gian", "Nguyên nhân", "Đã từng thành công"],
        [[
            a["problem"], a["batch"], f"{num(a['duration_seconds'])} giây", a["failure_reason"],
            "Có" if a["problem"] in e["latest_success"] else "Chưa",
        ] for a in latest_failures],
        widths=[2.4, 3.1, 2.3, 6.1, 2.3],
        font_size=8.4,
        alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER],
    )
    doc.add_paragraph(
        "cf1003a và cf1008a từng có report thành công ở batch hoàn tất trước đó, nên phụ lục vẫn dùng số liệu thành công gần nhất của hai bài này. "
        f"Năm bài chưa từng có report hoàn chỉnh trong phạm vi hợp lệ là {', '.join(s['never_success'])}."
    )

    doc.add_heading("7.3 Giới hạn của kết quả hiện tại", level=2)
    add_bullets(doc, [
        "TPR và TNR dựa trên nhãn lịch sử. Các nhãn chưa được xác minh lại bằng một judge chính thức cùng toolchain và checker.",
        "Reference consensus làm giảm rủi ro nhưng không chứng minh expected output đúng nếu các reference cùng chứa một lỗi.",
        "Ba mươi ba trên sáu mươi run thành công gần nhất dừng vì hết vòng, cho thấy điều kiện dừng sớm chưa đạt đồng đều.",
        "Mười hai run chưa phủ 100 phần trăm obligation của test model; coverage cũng phụ thuộc chất lượng model phân tích đề.",
        "Chỉ năm run có holdout metric. Phần lớn đánh giá tổng quát dựa vào oracle ngoài, chưa có holdout LLM đồng đều cho mọi bài.",
        "Native judge dành cho smoke test và không có mức cách ly như Docker. Metric cuối không nên lấy từ backend này cho source không tin cậy.",
        "Báo cáo không chạy lại code hoặc test, nên không xác nhận trạng thái thực thi của Docker, SSH hoặc LLM tại thời điểm viết.",
        "Workspace không có Git metadata khả dụng; báo cáo không thể gắn kết quả với commit cụ thể."
    ])

    doc.add_heading("8 Phần việc đã hoàn thành và hướng tiếp theo", level=1)
    doc.add_heading("8.1 Phần việc đã hoàn thành", level=2)
    add_bullets(doc, [
        "Hoàn thiện pipeline một lệnh cho một bài và batch tuần tự nhiều bài.",
        "Tách abstraction cho generator, judge, checker và oracle; runtime và sandbox đọc từ cấu hình.",
        "Bổ sung nhiều backend LLM, fallback theo nhiệm vụ và lưu raw response để audit.",
        "Chứng nhận reference, mutant và testcase trước khi tính metric.",
        "Tách development, holdout và oracle; khóa test suite trước đánh giá ngoài.",
        "Hỗ trợ Docker local, remote judge qua KHOI và native smoke test có cảnh báo.",
        "Tạo 72 bộ đề chuẩn hóa và 64 oracle bundle, tổng cộng 2.519 source có hash.",
        "Ghi report, iteration snapshot, execution log, latency và batch report để truy vết.",
        f"Sinh report hoàn chỉnh cho {s['ever_success_count']} bài và đo đủ oracle TPR/TNR cho {len(s['measured'])} bài.",
    ])
    doc.add_heading("8.2 Ưu tiên kỹ thuật tiếp theo", level=2)
    add_numbered(doc, [
        "Điều tra các bài TPR thấp bằng cách đối chiếu từng AC fail với checker, expected output, runtime và nhãn gốc.",
        "Tăng khả năng sinh validator và reference hợp lệ cho năm bài chưa từng hoàn tất; bổ sung prompt hoặc validator kiểm toán theo bài nếu cần.",
        "Cải thiện TNR cho các bài có WA sống sót bằng fault model cụ thể hơn và testcase nhắm survivor, sau đó đánh giá lại trên oracle khóa kín.",
        "Mở rộng holdout để có chỉ số tổng quát nội bộ trên nhiều bài hơn, đồng thời giữ nguyên ranh giới không đưa holdout vào feedback.",
        "Phân tích các run vượt 180 giây theo latency component để xác định phần LLM, chứng nhận mutant hay judge chiếm thời gian.",
        "Gắn mỗi đợt thực nghiệm với commit hoặc snapshot code để bảo đảm tái lập kết quả.",
        "Khi chuẩn bị số liệu cho luận văn, xác minh lại nhãn ngoại lệ và tách WA logic, RE và TLE trong bảng kết quả cuối."
    ])

    doc.add_heading("9 Kết luận tiến độ", level=1)
    doc.add_paragraph(
        "Dự án đã vượt qua giai đoạn chứng minh ý tưởng: pipeline hiện có thể phân tích đề, sinh và chứng nhận testcase, khóa bộ test, "
        "chấm source ngoài và lưu bằng chứng đầy đủ cho hàng chục bài. Kết quả định lượng cho thấy khả năng bảo toàn AC và bắt WA ở mức khá, "
        "với micro TPR và TNR đều trên 86 phần trăm. Điểm còn thiếu để chuyển thành kết quả luận văn vững là tính ổn định giữa các bài, "
        "quy trình điều tra false positive, mở rộng holdout và gắn kết quả với phiên bản code có thể tái lập."
    )

    doc.add_section(WD_SECTION.NEW_PAGE)
    appendix_section = doc.sections[-1]
    appendix_section.orientation = WD_ORIENT.PORTRAIT
    appendix_section.top_margin = Cm(1.7)
    appendix_section.bottom_margin = Cm(1.7)
    appendix_section.left_margin = Cm(1.4)
    appendix_section.right_margin = Cm(1.4)
    appendix_section.different_first_page_header_footer = False
    appendix_section.header.is_linked_to_previous = False
    appendix_section.footer.is_linked_to_previous = False
    set_repeat_header(appendix_section, "Phụ lục kết quả theo từng bài")
    add_page_number(appendix_section.footer.paragraphs[0])
    doc.add_heading("Phụ lục A Kết quả thành công gần nhất theo từng bài", level=1)
    doc.add_paragraph(
        "Mỗi dòng là run thành công gần nhất sau khi loại batch chưa hoàn tất. TPR và TNR là oracle metric theo nhãn. "
        "Dấu gạch ngang nghĩa là không có mẫu số đánh giá được."
    )
    problem_rows = []
    for r in sorted(s["runs"], key=lambda x: x["problem"].lower()):
        meta = e["metadata"].get(r["problem"], {})
        dataset = meta.get("dataset") or ("CodeNet" if r["problem"].startswith("p") else "Codeforces" if r["problem"].lower().startswith("cf") else "Khác")
        problem_rows.append([
            r["problem"], dataset, r["run"], r["cases"], r["iterations"], pct(r["coverage"]),
            pct(r["oracle_tpr"]), pct(r["oracle_tnr"]), classify(r["oracle_tpr"], r["oracle_tnr"]),
            num(r["generation_seconds"]),
        ])
    for start in range(0, len(problem_rows), 20):
        if start:
            doc.add_page_break()
        add_table(
            doc,
            ["Bài", "Nguồn", "Run", "Test", "Vòng", "Coverage", "TPR", "TNR", "Phân loại", "Giây"],
            problem_rows[start:start + 20],
            widths=[1.7, 1.8, 3.0, 1.1, 1.0, 1.7, 1.4, 1.4, 2.7, 1.4],
            font_size=6.7,
            alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT] + [WD_ALIGN_PARAGRAPH.CENTER] * 7,
        )

    doc.add_section(WD_SECTION.NEW_PAGE)
    source_section = doc.sections[-1]
    source_section.orientation = WD_ORIENT.PORTRAIT
    source_section.top_margin = Cm(1.7)
    source_section.bottom_margin = Cm(1.7)
    source_section.left_margin = Cm(1.4)
    source_section.right_margin = Cm(1.4)
    source_section.different_first_page_header_footer = False
    source_section.header.is_linked_to_previous = False
    source_section.footer.is_linked_to_previous = False
    set_repeat_header(source_section, "Phụ lục nguồn artifact")
    add_page_number(source_section.footer.paragraphs[0])
    doc.add_heading("Phụ lục B Nguồn artifact quan trọng", level=1)
    artifact_rows = [
        ["Mô tả hệ thống", "README.md, Architect.md, Specs.md"],
        ["Orchestrator", "testgen/pipeline.py"],
        ["Bộ sinh LLM", "testgen/ollama.py, testgen/claude.py, testgen/chatgpt.py"],
        ["Judge", "testgen/docker_judge.py, testgen/remote_judge.py, testgen/native_judge.py"],
        ["Checker", "testgen/checker.py, testgen/checkers/"],
        ["Cấu hình", "config/pipeline.json, config/judge.json, config/sandbox.json, config/runtimes.json"],
        ["Dữ liệu", "data/problems/<id>/statement.md, checker.json, oracle/manifest.json"],
        ["Kết quả một run", "output/.../report.json, test-suite.lock.json, iteration-*.json, executions-*.jsonl"],
        ["Kết quả batch", "output/batches/<timestamp>/batch-report.json và logs/"],
        ["Kiểm tra checker", e["validation_path"]],
    ]
    add_table(
        doc,
        ["Nhóm", "Đường dẫn"],
        artifact_rows,
        widths=[4.0, 14.0],
        font_size=8.3,
        alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT],
    )

    core = doc.core_properties
    core.title = "Báo cáo tiến độ hệ thống sinh testcase tự động"
    core.subject = "Final Testcase Generator"
    core.author = "Nhóm thực hiện khóa luận"
    core.keywords = "testcase generation, LLM, CodeNet, Codeforces, TPR, TNR"
    core.comments = "Tổng hợp từ code, dữ liệu và log hiện có đến ngày 12 tháng 10 năm 2026."
    doc.save(OUTPUT)


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    evidence = load_evidence()
    summary = summarize(evidence)
    charts = make_charts(evidence, summary)
    snapshot = {
        "captured_at": "2026-10-12T17:07:20+07:00",
        "summary": {k: v for k, v in summary.items() if k not in {"runs", "measured", "generation", "complete_batches", "excluded_batches", "latest_attempts"}},
        "latest_success": summary["runs"],
        "latest_attempts": summary["latest_attempts"],
        "batches": evidence["batches"],
    }
    (WORK_DIR / "evidence.json").write_text(json.dumps(snapshot, ensure_ascii=False, indent=2), encoding="utf-8")
    build_document(evidence, summary, charts)
    print(json.dumps({
        "output": str(OUTPUT),
        "size": OUTPUT.stat().st_size,
        "runs": summary["ever_success_count"],
        "measured": len(summary["measured"]),
        "pages_expected": "approximately 20 to 30",
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
